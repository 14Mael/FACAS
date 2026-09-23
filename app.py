from __future__ import annotations

import argparse
import base64
import json
import os
import re
import socket
import subprocess
import sys
import time
import traceback
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path


# PySide6 的 Qt DLL 在 PyInstaller 环境中分布在多个目录。
# 只注册 PySide6 和 shiboken6 自身目录，避免把其他依赖（尤其是
# Poppler 携带的 icuuc.dll）放到 Qt 的 DLL 搜索路径前面。
_QT_DLL_HANDLES = []


def prepare_frozen_dll_search_path() -> None:
    if not getattr(sys, "frozen", False):
        return
    runtime_root = Path(getattr(sys, "_MEIPASS", Path(sys.executable).resolve().parent))
    search_dirs = [runtime_root / "PySide6", runtime_root / "shiboken6"]

    # PyInstaller 可能会从其他包的依赖中收集同名 ICU DLL。它们不是
    # Qt 所需的 ICU 版本，不能把运行根目录加入 DLL 搜索路径，否则会
    # 在导入 QtCore 时触发 WinError 127。
    bundled_icu = {"icuuc.dll", "icudt78.dll"}
    has_conflicting_icu = any((runtime_root / name).is_file() for name in bundled_icu)
    if not has_conflicting_icu:
        root_text = str(runtime_root)
        path_entries = [entry for entry in os.environ.get("PATH", "").split(os.pathsep) if entry]
        if root_text not in path_entries:
            os.environ["PATH"] = os.pathsep.join([root_text, *path_entries])

    add_dll_directory = getattr(os, "add_dll_directory", None)
    if add_dll_directory is not None:
        for path in search_dirs:
            if path.is_dir():
                try:
                    _QT_DLL_HANDLES.append(add_dll_directory(str(path)))
                except OSError:
                    pass


prepare_frozen_dll_search_path()

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from playwright.sync_api import Error as PlaywrightError, Page, TimeoutError as PlaywrightTimeoutError, sync_playwright
from PySide6.QtCore import QDate, QObject, QThread, QTimer, Qt, Signal, Slot
from PySide6.QtGui import QColor, QFont, QPainter, QPen
from PySide6.QtWidgets import (
    QApplication,
    QCalendarWidget,
    QCheckBox,
    QDateEdit,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QScrollBar,
    QSizePolicy,
    QVBoxLayout,
    QWidget,
)

from bank_entry import MainWindow as BankEntryWindow


MODULES = {
    "现场销售": "bfc_sale_local_desk",
    "库房销售": "bfc_sale_desk",
    "拆解料销售": "bfc_sale_cjl_desk",
}

FORM_MODULES = {
    "综合付款单": "综合付款单",
    "车辆付款单": "车辆付款单",
    "采购单": "采购单",
}

ERP_URL = "https://erp.bfcgj.com/module.jsp?module=desk_main"
DEBUG_PORT = 9222
APP_VERSION = "v1.0.0"


def application_dir() -> Path:
    """Return the script folder, or the executable folder when packaged."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


class LoginRequired(RuntimeError):
    """Normal first-run state: the dedicated ERP browser needs a login."""


@dataclass
class SaleRow:
    sale_type: str
    sale_no: str
    sale_date: str
    payment_date: str
    material: str
    recovery_no: str
    total_amount: str
    tax: str
    total_weight: str
    remark: str


class DateEdit(QDateEdit):
    """带可见日历图标的日期输入框，避免系统箭头在主题中不可见。"""

    def paintEvent(self, event) -> None:
        super().paintEvent(event)
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        enabled = self.isEnabled()
        button_rect = self.rect().adjusted(self.width() - 31, 2, -2, -2)
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(QColor("#e5f2fc" if enabled else "#edf2f7"))
        painter.drawRoundedRect(button_rect, 7, 7)

        icon_color = QColor("#176fa8" if enabled else "#9aabba")
        painter.setPen(QPen(icon_color, 1.4))
        painter.setBrush(Qt.BrushStyle.NoBrush)
        icon_left = self.width() - 22
        icon_top = max(9, (self.height() - 13) // 2)
        painter.drawRoundedRect(icon_left, icon_top, 14, 12, 2, 2)
        painter.drawLine(icon_left, icon_top + 4, icon_left + 14, icon_top + 4)
        painter.drawLine(icon_left + 4, icon_top - 2, icon_left + 4, icon_top + 2)
        painter.drawLine(icon_left + 10, icon_top - 2, icon_left + 10, icon_top + 2)


class LogScrollBar(QScrollBar):
    """日志区域专用滚动条，统一绘制上下按钮，避免系统主题覆盖箭头图层。"""

    def paintEvent(self, event) -> None:
        super().paintEvent(event)
        if self.orientation() != Qt.Orientation.Vertical or self.height() < 36:
            return
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        button_height = 14
        button_color = QColor("#dcebf7")
        button_border = QColor("#c7dceb")
        arrow_color = QColor("#4f7fa2")
        width = self.width()
        painter.setPen(QPen(button_border, 1))
        painter.setBrush(button_color)
        painter.drawRoundedRect(1, 1, width - 2, button_height, 4, 4)
        painter.drawRoundedRect(1, self.height() - button_height - 1, width - 2, button_height, 4, 4)
        painter.setPen(QPen(arrow_color, 1.5, Qt.PenStyle.SolidLine, Qt.PenCapStyle.RoundCap, Qt.PenJoinStyle.RoundJoin))
        center_x = width // 2
        top_center = button_height // 2
        bottom_center = self.height() - button_height // 2 - 1
        painter.drawLine(center_x - 3, top_center + 1, center_x, top_center - 2)
        painter.drawLine(center_x, top_center - 2, center_x + 3, top_center + 1)
        painter.drawLine(center_x - 3, bottom_center - 1, center_x, bottom_center + 2)
        painter.drawLine(center_x, bottom_center + 2, center_x + 3, bottom_center - 1)


def clean(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def set_filter_checkbox(checkbox, checked: bool) -> None:
    """设置 ERP 筛选复选框，页面遮罩存在时退回到 DOM 事件。"""
    if checkbox.is_checked() == checked:
        return
    try:
        # ERP 筛选浮层偶尔会拦截鼠标事件，强制点击可绕过透明遮罩。
        checkbox.click(force=True, timeout=3000)
    except PlaywrightError:
        # 页面动画或遮罩持续存在时，直接触发原生 input/change 事件。
        checkbox.evaluate(
            """(el, target) => {
                const setter = Object.getOwnPropertyDescriptor(
                    HTMLInputElement.prototype, 'checked'
                ).set;
                setter.call(el, target);
                el.dispatchEvent(new Event('input', {bubbles: true}));
                el.dispatchEvent(new Event('change', {bubbles: true}));
            }""",
            checked,
        )
    if checkbox.is_checked() != checked:
        raise RuntimeError("ERP 筛选条件无法切换")


def safe_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*]', "_", clean(value))


def open_output_directory(path: Path) -> None:
    """Open a generated output directory in Windows Explorer."""
    directory = path if path.is_dir() else path.parent
    if directory.exists():
        os.startfile(str(directory))


def start_debug_edge() -> None:
    """Open a visible Edge window with a persistent automation profile."""
    edge_paths = (
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    )
    edge = next((item for item in edge_paths if item.exists()), None)
    if edge is None:
        raise RuntimeError("未找到 Microsoft Edge，请安装 Edge 后重试")
    profile = application_dir() / ".facas-edge-profile"
    profile.mkdir(exist_ok=True)
    subprocess.Popen([
        str(edge),
        f"--remote-debugging-port={DEBUG_PORT}",
        f"--user-data-dir={profile}",
        "--new-window",
        ERP_URL,
    ])


def debug_port_available() -> bool:
    """快速检测专用 Edge 端口，避免在界面线程中长时间阻塞。"""
    client = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        client.settimeout(0.08)
        return client.connect_ex(("127.0.0.1", DEBUG_PORT)) == 0
    except OSError:
        return False
    finally:
        client.close()


def live_erp_page(browser):
    pages = [
        page for context in browser.contexts for page in context.pages
        if not page.is_closed() and "erp.bfcgj.com" in page.url
    ]
    if not pages:
        raise RuntimeError("未找到仍然打开的 ERP 页面，请保持专用 Edge 和 ERP 页面打开")
    return pages[-1]


def save_page_pdf_fallback(page: Page, path: Path, a5_landscape: bool = False) -> None:
    """使用 Chromium CDP 保存 PDF，支持默认 A4 竖版和 A5 横版。"""
    session = page.context.new_cdp_session(page)
    paper_width, paper_height = (8.27, 5.83) if a5_landscape else (8.27, 11.69)
    result = session.send("Page.printToPDF", {
        "printBackground": True,
        # A5 横版时关闭 CSS 纸张优先，确保用户选择的纸张尺寸真正生效。
        "preferCSSPageSize": not a5_landscape,
        "paperWidth": paper_width,
        "paperHeight": paper_height,
        "marginTop": 0.2,
        "marginBottom": 0.2,
        "marginLeft": 0.2,
        "marginRight": 0.2,
    })
    path.write_bytes(base64.b64decode(result["data"]))


def find_print_markup(value):
    """Find HTML-like print markup in a PrintBill response recursively."""
    if isinstance(value, str):
        text = value.strip()
        if "<html" in text.lower() or "<body" in text.lower() or "lodop" in text.lower():
            return text
    if isinstance(value, dict):
        # PrintBill returns the template under anys.<module>.billprint.html;
        # it is a document fragment rather than a complete HTML document.
        for key, item in value.items():
            if key.lower() == "billprint.html" and isinstance(item, str):
                return item
            found = find_print_markup(item)
            if found:
                return found
    if isinstance(value, list):
        for item in value:
            found = find_print_markup(item)
            if found:
                return found
    return None


def page_print_markup(page: Page, expected_values: list[str] | None = None) -> str | None:
    """Read the bill HTML populated by plugInPrint.html after PrintBill."""
    # plugInPrint.html can be a hidden iframe or a separate context page.
    expected = [clean(value) for value in (expected_values or []) if clean(value) and len(clean(value)) >= 2]
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        best_markup = None
        best_score = -1
        for candidate_page in page.context.pages:
            for frame in candidate_page.frames:
                try:
                    bill = frame.locator("#bill1")
                    if not bill.count():
                        continue
                    body = bill.inner_html().strip()
                    if not body:
                        continue
                    text = clean(bill.inner_text())
                    score = sum(1 for value in expected if value in body or value in text)
                    if score > best_score:
                        style = frame.locator("#billStyle")
                        css = style.inner_text() if style.count() else ""
                        best_markup = f"<html><head><style>{css}</style></head><body>{body}</body></html>"
                        best_score = score
                except Exception:
                    continue
        # Never accept a zero-match frame: it may be the hidden print frame
        # left behind by the previous bill.
        if best_markup is not None and expected and best_score > 0:
            return best_markup
        time.sleep(0.2)
    return None


def print_payload_values(value) -> list[str]:
    values = []
    if isinstance(value, dict):
        for item in value.values():
            values.extend(print_payload_values(item))
    elif isinstance(value, list):
        for item in value:
            values.extend(print_payload_values(item))
    elif isinstance(value, (str, int, float)):
        text = clean(str(value))
        if len(text) >= 2 and not text.startswith("data:"):
            values.append(text)
    return values


def dismiss_erp_print_error(page: Page, log) -> None:
    """Close the ERP's DOM error dialog raised when LODOP is unavailable."""
    try:
        dialog = page.locator(".yc-view-message:visible").last
        if not dialog.count():
            return
        message = clean(dialog.inner_text())
        if message:
            log(f"ERP 打印插件提示，已关闭: {message}")
        confirm = dialog.get_by_role("button", name="确认")
        if confirm.count() and confirm.is_visible():
            confirm.click(timeout=2000)
        else:
            dialog.locator("button").last.click(timeout=2000)
    except (PlaywrightError, PlaywrightTimeoutError):
        pass


def ensure_clodop(page: Page, log) -> bool:
    """Load the local C-Lodop bridge when ERP did not include it in the page."""
    try:
        ready = page.evaluate("""() => {
            const printer = window.LODOP || window.CLODOP;
            return !!(printer && typeof printer.CVERSION === 'function'
                && typeof printer.SET_LICENSES === 'function');
        }""")
        if ready:
            return True
        page.evaluate("""() => new Promise((resolve, reject) => {
            const printer = window.LODOP || window.CLODOP;
            if (printer && typeof printer.CVERSION === 'function'
                    && typeof printer.SET_LICENSES === 'function') { resolve(true); return; }
            const script = document.createElement('script');
            script.src = 'http://localhost:8000/CLodopfuncs.js';
            script.onload = () => resolve(true);
            script.onerror = () => reject(new Error('无法加载本机 C-Lodop 脚本'));
            document.head.appendChild(script);
        })""")
        page.wait_for_function("() => { const printer = window.LODOP || window.CLODOP; return !!(printer && typeof printer.CVERSION === 'function' && typeof printer.SET_LICENSES === 'function'); }", timeout=5000)
        log("已加载本机 C-Lodop 打印组件")
        return True
    except Exception as exc:
        log(f"C-Lodop 未就绪: {exc}")
        return False


def visible_print_preview_button(page: Page):
    """查找当前详情页可见的打印预览按钮，避免误取隐藏按钮。"""
    candidates = [
        page.get_by_role("button", name=re.compile(r"打印\s*预览")),
        page.locator("button").filter(has_text=re.compile(r"打印\s*预览")),
        page.get_by_text("打印预览", exact=True),
    ]
    for locator in candidates:
        try:
            for index in range(locator.count() - 1, -1, -1):
                candidate = locator.nth(index)
                if candidate.is_visible() and candidate.is_enabled():
                    return candidate
        except PlaywrightError:
            continue
    return None


def save_printbill_pdf(page: Page, path: Path, log, expected_values: list[str] | None = None,
                       a5_landscape: bool = False) -> bool:
    """Capture the ERP PrintBill response and render returned markup."""
    # 防止同名旧文件在本次打印失败时被误判为新生成文件。
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.unlink(missing_ok=True)
    except OSError as exc:
        log(f"清理旧 PDF 失败: {path.name} ({exc})")
    # Do not click "单据打印": that action invokes the native Windows print
    # dialog. Preview is the non-destructive route used for PDF extraction.
    print_button = visible_print_preview_button(page)
    if print_button is None:
        log(f"当前详情页未找到“打印预览”按钮: {path.name}")
        return False
    clodop_ready = ensure_clodop(page, log)
    if not clodop_ready:
        # PrintBill 响应和渲染的 HTML 有时仍可正常返回，不因 ERP 打印插件
        # 弹窗直接跳过，后面会捕获并关闭该弹窗后继续生成 PDF。
        log(f"LODOP/C-Lodop 未完全就绪，继续尝试读取 PrintBill: {path.name}")
    try:
        with page.expect_response(lambda response: "_name=PrintBill" in response.url, timeout=10000) as response_info:
            print_button.click(timeout=5000, no_wait_after=True)
        response = response_info.value
        if not page.is_closed():
            page.wait_for_timeout(500)
            dismiss_erp_print_error(page, log)
        body = response.text()
        try:
            payload = response.json()
        except Exception:
            payload = body
        payload_values = print_payload_values(payload)
        payload_business_values = sorted({
            value for value in payload_values
            if re.search(r"\$[A-Z0-9]{6,}", value)
        }, key=len, reverse=True)
        # ERP may return a fixed template without the internal bill number.
        # In that case validate the rendered frame against identifiers from
        # the detail page that is currently open.
        business_values = []
        for value in [*(expected_values or []), *payload_business_values]:
            value = clean(str(value or ""))
            if value and len(value) >= 2 and value not in business_values:
                business_values.append(value)
        if not payload_business_values:
            log("PrintBill 响应未包含内部单号，改用当前单据页面标识校验")
        # The response often contains only the fixed template. Prefer the
        # fully populated #bill1 rendered by the ERP print iframe.
        markup = page_print_markup(page, business_values)
        if not markup:
            response_markup = find_print_markup(payload)
            # PrintBill 可能返回固定模板，业务值位于同级 data/vars 中，
            # 不一定直接出现在模板 HTML 里。响应刚刚由当前单据触发，
            # 因此只要模板具有可渲染结构即可作为本次打印模板。
            if response_markup and len(response_markup) >= 200:
                markup = response_markup
                log("使用本次 PrintBill 响应中的固定打印模板")
        if not markup:
            debug_path = path.with_suffix(".print-response.txt")
            debug_path.write_text(body, encoding="utf-8")
            html_path = path.with_suffix(".print-page.html")
            html_path.write_text(page.content(), encoding="utf-8")
            log(f"PrintBill 已响应，但未发现可渲染 HTML: {body[:200]}")
            log(f"已保存打印响应和打印页面供排查: {debug_path.name}, {html_path.name}")
            return False
        preview = page.context.new_page()
        try:
            preview.set_content(markup, wait_until="networkidle")
            save_page_pdf_fallback(preview, path, a5_landscape=a5_landscape)
        finally:
            preview.close()
        return path.is_file() and path.stat().st_size > 0
    except Exception as exc:
        log(f"捕获 PrintBill 失败: {exc}")
        return False


def data_grid(page: Page, expected_markers: tuple[str, ...] = ()):
    """Return the visible ERP list grid matching its business headers."""
    grids = page.locator("table.yc-view-grid-table")
    header_markers = ("销售单号", "票据号码", "单据日期", "客户名称", "对方单位")
    candidates: list[tuple[int, object]] = []
    for i in range(grids.count()):
        grid = grids.nth(i)
        try:
            if not grid.is_visible():
                continue
            text = clean(grid.inner_text())
            marker_count = sum(marker in text for marker in header_markers)
            expected_count = sum(marker in text for marker in expected_markers)
            record_count = len(re.findall(r"\$[A-Z]{2}\d{6,}|20\d{2}-\d{2}-\d{2}", text))
            data_rows = grid.locator("tbody tr").count() or grid.locator("tr").count()
            has_action = grid.get_by_text("查看", exact=True).count() > 0
            subtotal_penalty = 20 if any(label in text for label in ("本页小计", "合计")) else 0
            if data_rows and (marker_count or record_count):
                # ERP separates the header and body into different tables.
                # Real records must outweigh header-marker matches, otherwise
                # an empty header table wins and detail rows cannot be opened.
                score = (record_count * 50 + expected_count * 20
                         + marker_count * 5 + data_rows + int(has_action) * 3
                         - subtotal_penalty)
                candidates.append((score, grid))
        except PlaywrightError:
            continue
    if candidates:
        return max(candidates, key=lambda item: item[0])[1]
    expected_text = "、".join(expected_markers) if expected_markers else "销售单号或票据号码"
    raise RuntimeError(f"未识别到包含{expected_text}的数据表格")


def click_visible_menu_item(page: Page, name: str) -> None:
    # The ERP may render the navigation tree inside an iframe and may wrap the
    # label in a non-clickable span. Search every frame and click its nearest
    # clickable ancestor, preferring the leftmost visible match.
    visible_candidates = []
    for frame in page.frames:
        candidates = frame.get_by_text(name, exact=True)
        for index in range(candidates.count()):
            candidate = candidates.nth(index)
            try:
                if not candidate.is_visible():
                    continue
                clickable = candidate.locator("xpath=ancestor-or-self::*[self::button or self::a or @role='menuitem' or contains(@class,'menu')][1]")
                target = clickable if clickable.count() else candidate
                box = target.bounding_box()
                if box:
                    visible_candidates.append((box["x"], box["y"], target))
            except PlaywrightError:
                continue
    for _, _, candidate in sorted(visible_candidates, key=lambda item: (item[0], item[1])):
        try:
            candidate.scroll_into_view_if_needed(timeout=3000)
            candidate.click(timeout=10000, force=True)
            page.wait_for_timeout(500)
            return
        except PlaywrightError:
            continue
    raise RuntimeError(f"未找到可点击的菜单项: {name}")


def expand_menu_section(page: Page, name: str, child_name: str | None = None) -> None:
    """Expand a first-level ERP menu without collapsing an already-open one."""
    for frame in page.frames:
        labels = frame.get_by_text(name, exact=True)
        for index in range(labels.count()):
            label = labels.nth(index)
            try:
                if not label.is_visible():
                    continue
                target = label.locator("xpath=ancestor-or-self::*[self::button or self::a or @role='treeitem' or @role='menuitem' or contains(@class,'menu')][1]")
                if not target.count():
                    target = label
                if child_name:
                    child = frame.get_by_text(child_name, exact=True)
                    if any(child.nth(i).is_visible() for i in range(child.count())):
                        return
                expanded = target.get_attribute("aria-expanded")
                classes = (target.get_attribute("class") or "").lower()
                # aria-expanded is authoritative when provided. For older ERP
                # markup, infer collapsed state from common class names.
                collapsed = expanded == "false" or any(token in classes for token in ("collapsed", "closed", "fold"))
                if expanded is None and not collapsed:
                    target.click(timeout=8000, force=True)
                    page.wait_for_timeout(400)
                    return
                if collapsed:
                    target.click(timeout=8000, force=True)
                    page.wait_for_timeout(400)
                return
            except PlaywrightError:
                continue
    raise RuntimeError(f"未找到一级菜单或无法展开: {name}")


def apply_date_filters(filter_table, start: str, end: str, before_search=None) -> bool:
    """Apply the first/last date fields without assuming mixed input indexes."""
    checkboxes = filter_table.locator("input[type='checkbox']:visible")
    text_candidates = filter_table.locator("input:not([type='checkbox']):visible")
    editable_inputs = []
    for index in range(text_candidates.count()):
        candidate = text_candidates.nth(index)
        try:
            if candidate.is_editable():
                editable_inputs.append(candidate)
        except PlaywrightError:
            continue
    if checkboxes.count() < 4 or len(editable_inputs) < 2:
        return False
    # Resolve each checkbox by the label in its nearest filter container,
    # rather than by position (the ERP changes order between modules).
    date_controls = []
    for index in range(checkboxes.count()):
        checkbox = checkboxes.nth(index)
        label_text = checkbox.evaluate("""el => {
            let node = el;
            for (let i = 0; i < 6 && node; i++, node = node.parentElement) {
                const text = (node.innerText || '').replace(/\\s+/g, '');
                if (text.includes('销售日期从') || text.includes('销售日期到')) return text;
            }
            return '';
        }""")
        if "销售日期从" in label_text or "销售日期到" in label_text:
            # A filter field is rendered inside one table cell: checkbox,
            # label, and input live in the same td. Keep the lookup within
            # that exact HTML cell so it cannot select 客户名称 or 回收单号.
            container = checkbox.locator("xpath=ancestor::td[1]")
            date_input = container.locator("input:not([type='checkbox']):visible").first
            if date_input.count():
                value = start if "销售日期从" in label_text else end
                date_controls.append((checkbox, date_input, value))
    if len(date_controls) != 2:
        date_controls = []
        for label, value in (("销售日期从", start), ("销售日期到", end)):
            label_locator = filter_table.get_by_text(label, exact=False).filter(visible=True)
            found = False
            for index in range(label_locator.count()):
                label_node = label_locator.nth(index)
                cell = label_node.locator("xpath=ancestor::td[1]")
                checkbox = cell.locator("input[type='checkbox']:visible").first
                date_input = cell.locator("input:not([type='checkbox']):visible").first
                if checkbox.count() and date_input.count():
                    date_controls.append((checkbox, date_input, value))
                    found = True
                    break
            if not found:
                return False
    for index in range(checkboxes.count()):
        checkbox = checkboxes.nth(index)
        if checkbox.is_checked():
            set_filter_checkbox(checkbox, False)
    for checkbox, date_input, value in date_controls:
        set_filter_checkbox(checkbox, True)
        date_input.fill(value)
        date_input.press("Enter")
    search = filter_table.get_by_role("button", name="搜索")
    if search.count():
        if before_search is not None:
            before_search()
        search.click()
        return True
    return False


def apply_voucher_filters(page: Page, filter_table, start: str, end: str, before_search=None) -> bool:
    """按凭证探测器筛选控件属性设置日期并包含已入账。"""
    def control_by_filter(*names):
        for name in names:
            locator = filter_table.locator(f'input[filter-name="{name}"]:visible')
            if locator.count():
                return locator.first
        return None

    def date_control(*names):
        checkbox = control_by_filter(*names)
        if checkbox is None:
            return None
        container = checkbox.locator("xpath=ancestor::tr[1]")
        date_input = container.locator("input:not([type='checkbox']):visible").first
        if not date_input.count():
            container = checkbox.locator("xpath=ancestor::td[1]")
            date_input = container.locator("input:not([type='checkbox']):visible").first
        if not date_input.count() or not date_input.is_editable():
            return None
        return checkbox, date_input

    start_control = date_control("billDateFrom", "bill_datefrom")
    end_control = date_control("BillDateTo", "billDateTo", "bill_dateto")
    if start_control is None or end_control is None:
        return False

    include_checkbox = control_by_filter("accounted", "in_stock", "include_stock")
    if include_checkbox is None:
        labels = filter_table.get_by_text("包含已入账", exact=False)
        if not labels.count():
            labels = filter_table.get_by_text("包含已入库", exact=False)
        if labels.count():
            cell = labels.first.locator("xpath=ancestor::td[1]")
            candidate = cell.locator("input[type='checkbox']:visible").first
            if candidate.count():
                include_checkbox = candidate
    if include_checkbox is None:
        return False

    # 关闭其他筛选条件，只保留日期范围和包含已入账。
    keep = {"billDateFrom", "bill_datefrom", "BillDateTo", "billDateTo", "bill_dateto",
            "accounted", "in_stock", "include_stock"}
    for index in range(filter_table.locator("input[type='checkbox']:visible").count()):
        checkbox = filter_table.locator("input[type='checkbox']:visible").nth(index)
        filter_name = checkbox.get_attribute("filter-name") or ""
        if filter_name not in keep:
            set_filter_checkbox(checkbox, False)
    set_filter_checkbox(start_control[0], True)
    set_filter_checkbox(end_control[0], True)
    set_filter_checkbox(include_checkbox, True)
    start_control[1].fill(start)
    end_control[1].fill(end)
    start_control[1].press("Enter")
    end_control[1].press("Enter")

    submit = None
    for name in ("确定", "确认"):
        candidate = filter_table.get_by_role("button", name=name)
        if candidate.count() and candidate.is_visible():
            submit = candidate.last
            break
        candidate = page.get_by_role("button", name=name)
        if candidate.count() and candidate.is_visible():
            submit = candidate.last
            break
    if submit is None:
        submit = filter_table.get_by_role("button", name="搜索")
    if not submit.count():
        return False
    if before_search is not None:
        before_search()
    submit.click(timeout=5000, force=True)
    return True

def apply_sale_date_filters(filter_table, start: str, end: str, before_search=None) -> bool:
    """Use the stable sales-page checkbox/input pairs from the ERP layout."""
    inputs = filter_table.locator("input:visible")
    if inputs.count() < 8:
        return False
    # Disable every filter first, then enable only sales-date from/to.
    for index in range(inputs.count()):
        control = inputs.nth(index)
        if (control.get_attribute("type") or "").lower() == "checkbox" and control.is_checked():
            set_filter_checkbox(control, False)
    for checkbox_index, input_index, value in ((0, 1, start), (6, 7, end)):
        checkbox = inputs.nth(checkbox_index)
        date_input = inputs.nth(input_index)
        if not checkbox.is_checked():
            set_filter_checkbox(checkbox, True)
        date_input.fill(value)
        date_input.press("Enter")
    search = filter_table.get_by_role("button", name="搜索")
    if search.count():
        if before_search is not None:
            before_search()
        search.click()
        return True
    return False


def remove_response_listener(page: Page, callback) -> None:
    """Release a response callback even when the ERP page is closing."""
    try:
        page.remove_listener("response", callback)
    except PlaywrightError:
        pass


def find_named_blocks(value, name: str, module: str | None = None) -> list[dict]:
    """Return every named ERP response block, optionally for one module."""
    found: list[dict] = []
    if isinstance(value, dict):
        if (value.get("name") == name and isinstance(value.get("rowsData"), list)
                and (module is None or value.get("_amn") == module)):
            found.append(value)
        for child in value.values():
            found.extend(find_named_blocks(child, name, module))
    elif isinstance(value, list):
        for child in value:
            found.extend(find_named_blocks(child, name, module))
    return found


DETAIL_FIELD_MAP = {
    # Fields confirmed from the ERP's bfc_sale_local_bill data_d response.
    "现场销售": {
        "material": "c14", "recovery_no": "c3", "amount": "c12",
        "tax": "c17", "weight": "c16",
    },
    # Fields confirmed from the ERP's bfc_sale_bill data_d response.
    "库房销售": {
        "material": "c12", "recovery_no": None, "amount": "c10",
        "tax": "c15", "weight": "c26",
    },
    "拆解料销售": {
        "material": "c14", "recovery_no": None, "amount": "c12",
        "tax": "c17", "weight": "c16",
    },
}


def numeric_value(value) -> float | None:
    """Convert ERP numeric fields while treating blanks as absent values."""
    if value is None or clean(str(value)) == "":
        return None
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return None


def format_amount(value: float | None) -> str:
    return f"{value:,.2f}" if value is not None else ""


def erp_date(value) -> str:
    if isinstance(value, (int, float)):
        return datetime.fromtimestamp(value / 1000).strftime("%Y-%m-%d")
    return clean(str(value or ""))


SALE_LIST_FIELD_MAP = {
    "现场销售": {
        "sale_no": "c0", "sale_date": "c4", "payment_date": "c9",
        "customer": "c3", "amount": "c5", "weight": "c16",
    },
    "库房销售": {
        "sale_no": "c0", "sale_date": "c4", "payment_date": "c5",
        "customer": "c3", "amount": "c6", "weight": "c15",
    },
    "拆解料销售": {
        "sale_no": "c0", "sale_date": "c4", "payment_date": "c9",
        "customer": "c3", "amount": "c5", "weight": "c15",
    },
}


def sale_list_row(sale_type: str, row: dict) -> tuple[str, str, str, str, str, str]:
    """Map the ERP data_list response for the selected sales desk."""
    fields = SALE_LIST_FIELD_MAP[sale_type]
    return (
        clean(str(row.get(fields["sale_no"]) or "")),
        erp_date(row.get(fields["sale_date"])),
        erp_date(row.get(fields["payment_date"])),
        clean(str(row.get(fields["customer"]) or "")),
        format_amount(numeric_value(row.get(fields["amount"]))),
        format_amount(numeric_value(row.get(fields["weight"]))),
    )


def click_sale_detail(page: Page, sale_no: str) -> None:
    """Open a visible list row by the API-provided bill number."""
    matches = page.get_by_text(sale_no, exact=True)
    for index in range(matches.count()):
        bill_cell = matches.nth(index)
        try:
            if not bill_cell.is_visible():
                continue
            row = bill_cell.locator("xpath=ancestor::tr[1]")
            view = row.get_by_text("查看", exact=True)
            if view.count() and view.is_visible():
                view.click()
                return
        except PlaywrightError:
            continue
    raise RuntimeError(f"页面中未找到销售单 {sale_no} 的查看操作")


def detail_api_summary(sale_type: str, rows: list[dict]) -> tuple[str, str, str, str, str]:
    """Read the first item and aggregate numeric detail fields from data_d."""
    fields = DETAIL_FIELD_MAP[sale_type]
    material = ""
    recovery_no = ""
    amount_total = tax_total = weight_total = 0.0
    amount_seen = tax_seen = weight_seen = False
    for row in rows:
        if not material:
            material = clean(str(row.get(fields["material"]) or ""))
        if not recovery_no and fields["recovery_no"]:
            recovery_no = clean(str(row.get(fields["recovery_no"]) or ""))
        amount = numeric_value(row.get(fields["amount"]))
        tax = numeric_value(row.get(fields["tax"]))
        weight = numeric_value(row.get(fields["weight"]))
        if amount is not None:
            amount_total += amount
            amount_seen = True
        if tax is not None:
            tax_total += tax
            tax_seen = True
        if weight is not None:
            weight_total += weight
            weight_seen = True
    return (
        material,
        recovery_no,
        format_amount(amount_total if amount_seen else None),
        format_amount(tax_total if tax_seen else None),
        format_amount(weight_total if weight_seen else None),
    )


def sale_master_api_summary(rows: list[dict]) -> tuple[str, str, str]:
    """Read totals, tax total, and remark from a sales data_m response."""
    if not rows:
        return "", "", ""
    row = rows[0]
    return (
        format_amount(numeric_value(row.get("c13"))),
        format_amount(numeric_value(row.get("c14"))),
        clean(str(row.get("c9") or "")),
    )


def voucher_value(row: dict, header: str) -> str:
    header = clean(header)
    key_map = {
        "票据号码": "c0", "单据号码": "c0", "编号": "c0",
        "对方单位": "c9", "供应商": "c9", "客户": "c9",
        "单据日期": "c3", "完成日期": "c4", "日期": "c3",
        "金额": "c6", "结算方式": "c5", "制单": "c7",
    }
    key = next((source for label, source in key_map.items() if label in header), None)
    value = row.get(key, "") if key else ""
    if key in ("c3", "c4") and isinstance(value, (int, float)):
        return datetime.fromtimestamp(value / 1000).strftime("%Y-%m-%d")
    if value is None:
        return ""
    return clean(str(value).replace("\u00a0", " "))




def scrape_module(page: Page, sale_type: str, start: str, end: str, log,
                  pdf_dir: Path | None = None, generated_pdfs: set[Path] | None = None,
                  pdf_a5_landscape: bool = False) -> list[SaleRow]:
    if page.is_closed():
        raise RuntimeError("ERP 页面已关闭，请重新登录专用 Edge")
    log(f"{sale_type}: 正在打开销售列表")
    expand_menu_section(page, "销售管理", sale_type)
    click_visible_menu_item(page, sale_type)
    page.wait_for_timeout(500)
    data_grid(page, ("销售单号",)).wait_for(state="visible", timeout=15000)

    sale_page_rows: dict[int, list[dict]] = {}
    sale_page_batches: list[list[dict]] = []
    sale_page_signatures: set[tuple[str, ...]] = set()
    sale_total_pages = 0
    sale_total_rows = 0
    sale_page_size = 0
    capture_sale_enabled = False
    sale_response_seen = False

    def capture_sale_response(response):
        nonlocal sale_total_pages, sale_total_rows, sale_page_size, sale_response_seen
        if not capture_sale_enabled:
            return
        if "erp.bfcgj.com" not in response.url:
            return
        try:
            payload = response.json()
        except Exception:
            return
        for block in find_named_blocks(payload, "data_list", MODULES[sale_type]):
            sale_response_seen = True
            rows = list(block["rowsData"])
            signature = tuple(clean(str(row.get("c0") or "")) for row in rows)
            if signature and signature not in sale_page_signatures:
                sale_page_signatures.add(signature)
                sale_page_batches.append(rows)
            page_index = int(block.get("page") or 0)
            sale_page_rows[page_index] = rows
            sale_total_pages = max(sale_total_pages, int(block.get("pages") or 0))
            sale_total_rows = max(sale_total_rows, int(block.get("totalRows") or 0))
            sale_page_size = max(sale_page_size, int(block.get("pageRows") or len(rows)))

    def begin_sale_capture():
        nonlocal capture_sale_enabled, sale_response_seen, sale_total_pages, sale_total_rows, sale_page_size
        sale_page_rows.clear()
        sale_page_batches.clear()
        sale_page_signatures.clear()
        sale_total_pages = sale_total_rows = sale_page_size = 0
        sale_response_seen = False
        capture_sale_enabled = True

    page.on("response", capture_sale_response)
    output: list[SaleRow] = []
    detail_failures: list[str] = []
    listed_count = skipped_payment_count = page_no = 0
    seen_sale_nos: set[str] = set()
    try:
        filter_tables = page.locator("table.yc-view-free-table:visible")
        filter_table = filter_tables.first if filter_tables.count() else page.locator("table.yc-view-free-table").first
        # Ignore any delayed default-list response emitted while the desk was
        # opening. Only responses triggered by the requested filter/pagination
        # are allowed into the page buffers.
        sale_page_rows.clear()
        sale_page_batches.clear()
        sale_page_signatures.clear()
        sale_total_pages = sale_total_rows = sale_page_size = 0
        sale_response_seen = False
        capture_sale_enabled = False
        if not apply_sale_date_filters(filter_table, start, end, begin_sale_capture):
            raise RuntimeError(f"{sale_type}日期筛选未成功执行，已停止导出")
        log(f"{sale_type}: 已提交销售日期筛选 {start} 至 {end}")
        deadline = time.monotonic() + 10
        while not sale_response_seen and time.monotonic() < deadline:
            page.wait_for_timeout(200)
        if not sale_response_seen:
            raise RuntimeError(f"未捕获{sale_type} data_list 响应，已停止导出")
        if not sale_page_batches:
            if sale_total_rows:
                raise RuntimeError(f"{sale_type} data_list 返回空数据但 totalRows={sale_total_rows}")
            log(f"{sale_type}: 日期范围内未查询到销售单")
            return output
        sale_page_rows[0] = sale_page_batches[0]
        expected_pages = 0
        if sale_total_rows and sale_page_size:
            expected_pages = (sale_total_rows + sale_page_size - 1) // sale_page_size
        effective_total_pages = max(sale_total_pages, expected_pages)
        if sale_total_pages:
            log(f"{sale_type}: 接口共 {sale_total_pages} 页")
        if expected_pages and expected_pages != sale_total_pages:
            log(f"{sale_type}: 接口共 {sale_total_rows} 条，预计 {expected_pages} 页")

        while True:
            if page.is_closed():
                raise RuntimeError("ERP 页面已关闭，无法继续读取销售列表")
            page_no += 1
            log(f"{sale_type}: 正在读取第 {page_no} 页")
            current_rows = sale_page_rows.get(page_no - 1)
            if current_rows is None and page_no - 1 < len(sale_page_batches):
                current_rows = sale_page_batches[page_no - 1]
            if current_rows is None:
                raise RuntimeError(f"未捕获{sale_type}第 {page_no} 页 data_list 响应")
            for sale_no, sale_date, payment_date, _customer, _amount, _list_weight in (
                    sale_list_row(sale_type, row) for row in current_rows):
                if page.is_closed():
                    raise RuntimeError(
                        f"{sale_type} 读取销售单 {sale_no} 前 ERP 页面已关闭，请重新打开专用 Edge 后重试"
                    )
                if sale_no in seen_sale_nos:
                    continue
                seen_sale_nos.add(sale_no)
                listed_count += 1
                if not payment_date or not re.search(r"\d{4}-\d{2}-\d{2}", payment_date):
                    skipped_payment_count += 1
                    continue
                material = recovery_no = remark = tax = ""
                amount = weight = ""
                detail_succeeded = False
                close = None
                capture_detail_response = None
                try:
                    log(f"{sale_type}: 正在读取销售单 {sale_no} 详情")
                    detail_response_rows: list[dict] = []
                    master_response_rows: list[dict] = []
                    detail_row_keys: set[tuple] = set()
                    master_row_keys: set[tuple] = set()

                    def response_row_key(row: dict) -> tuple:
                        """Identify repeated ERP rows across duplicate responses."""
                        for field in ("__rowid", "__rownum"):
                            value = row.get(field)
                            if value is not None and clean(str(value)):
                                return field, str(value)
                        return "data", json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)

                    def capture_detail_response(response):
                        if "erp.bfcgj.com" not in response.url:
                            return
                        try:
                            payload = response.json()
                        except Exception:
                            return
                        detail_module = MODULES[sale_type].replace("_desk", "_bill")
                        for block in find_named_blocks(payload, "data_d", detail_module):
                            for row in block["rowsData"]:
                                if clean(str(row.get("c0") or "")) != sale_no:
                                    continue
                                key = response_row_key(row)
                                if key not in detail_row_keys:
                                    detail_row_keys.add(key)
                                    detail_response_rows.append(row)
                        for block in find_named_blocks(payload, "data_m", detail_module):
                            for row in block["rowsData"]:
                                if clean(str(row.get("c0") or "")) != sale_no:
                                    continue
                                key = response_row_key(row)
                                if key not in master_row_keys:
                                    master_row_keys.add(key)
                                    master_response_rows.append(row)

                    page.on("response", capture_detail_response)
                    click_sale_detail(page, sale_no)
                    close = page.get_by_role("button", name="关闭").last
                    close.wait_for(state="visible", timeout=10000)
                    deadline = time.monotonic() + 10
                    while (not detail_response_rows or not master_response_rows) and time.monotonic() < deadline:
                        page.wait_for_timeout(200)
                    if not detail_response_rows or not master_response_rows:
                        raise RuntimeError("未捕获销售详情完整 data_m/data_d 响应")
                    material, recovery_no, calculated_amount, detail_tax, weight = detail_api_summary(sale_type, detail_response_rows)
                    master_amount, master_tax, remark = sale_master_api_summary(master_response_rows)
                    amount = master_amount or calculated_amount
                    tax = master_tax if numeric_value(master_tax) not in (None, 0.0) else detail_tax
                    detail_succeeded = True
                    log(f"{sale_type}: 已读取销售单 {sale_no} 详情")
                    if pdf_dir is not None:
                        pdf_path = pdf_dir / f"{safe_filename(sale_type)}-{safe_filename(sale_date)}-{safe_filename(sale_no)}.pdf"
                        expected_print_values = [sale_no, sale_no.lstrip("$")[-8:], sale_date]
                        if save_printbill_pdf(
                                page, pdf_path, log, expected_print_values,
                                a5_landscape=pdf_a5_landscape):
                            if generated_pdfs is not None:
                                generated_pdfs.add(pdf_path.resolve())
                            log(f"已保存 PDF: {pdf_path.name}")
                        else:
                            log(f"未生成票据 PDF: {pdf_path.name}")
                except (RuntimeError, PlaywrightTimeoutError, PlaywrightError, IndexError, OSError) as exc:
                    detail_failures.append(sale_no)
                    log(f"详情读取失败，已跳过销售单 {sale_no}: {exc}")
                finally:
                    if capture_detail_response is not None:
                        remove_response_listener(page, capture_detail_response)
                    try:
                        if close is not None and not page.is_closed() and close.count() and close.is_visible():
                            close.click(timeout=3000, force=True, no_wait_after=True)
                            if not page.is_closed():
                                page.wait_for_timeout(300)
                    except PlaywrightError:
                        log(f"详情关闭失败: {sale_no}")
                if detail_succeeded:
                    output.append(SaleRow(sale_type, sale_no, sale_date, payment_date, material, recovery_no, amount, tax, weight, remark))

            # The API's page count is authoritative. Do not click the next
            # button after the final page, where no new data_list response can
            # arrive and the old timeout misleadingly reports a missing page.
            if effective_total_pages and page_no >= effective_total_pages:
                break
            target_page = str(page_no + 1)
            page_links = page.locator("li.paging-item:visible")
            next_buttons = page.locator('[paging-btn="nextPage"]:visible')
            advance = None
            for link_index in range(page_links.count()):
                candidate = page_links.nth(link_index)
                if clean(candidate.inner_text()) == target_page:
                    advance = candidate
                    break
            if advance is not None:
                pass
            else:
                if not next_buttons.count():
                    break
                advance = next_buttons.last
                if advance.is_disabled() or "disabled" in (advance.get_attribute("class") or ""):
                    break
            previous_batch_count = len(sale_page_batches)
            selected_page = ""
            selected_items = page.locator("li.paging-item.paging-item-selected:visible")
            if selected_items.count():
                selected_page = clean(selected_items.last.inner_text())
            already_selected = selected_page == target_page
            batch_ready = already_selected and len(sale_page_batches) > page_no
            if already_selected and len(sale_page_batches) > page_no:
                sale_page_rows[page_no] = sale_page_batches[page_no]
            elif not already_selected:
                advance.scroll_into_view_if_needed()
                advance.click(timeout=10000, force=True)
            deadline = time.monotonic() + 10
            if not batch_ready:
                while len(sale_page_batches) <= previous_batch_count and time.monotonic() < deadline:
                    page.wait_for_timeout(200)
                if len(sale_page_batches) <= previous_batch_count:
                    raise RuntimeError(f"未捕获{sale_type}第 {page_no + 1} 页 data_list 响应")
                sale_page_rows[page_no] = sale_page_batches[previous_batch_count]
    finally:
        remove_response_listener(page, capture_sale_response)
    failure_note = f"；详情失败 {len(detail_failures)} 张: {'、'.join(detail_failures)}" if detail_failures else ""
    log(f"{sale_type}: 列表销售单 {listed_count} 张，因收款日期为空跳过 {skipped_payment_count} 张，导出 {len(output)} 张{failure_note}")
    return output


def write_sheet(ws, title: str, headers: list[str], rows: list[list[str]], widths: list[int] | None = None) -> None:
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    if widths is not None:
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + index)].width = width
        return
    for column in ws.columns:
        letter = column[0].column_letter
        max_len = min(30, max([len(str(cell.value or "")) for cell in column] + [10]))
        ws.column_dimensions[letter].width = max_len + 2


def export_report(rows: list[SaleRow], sheets: dict[str, tuple[list[str], list[list[str]]]], path: Path) -> None:
    """Write only data selected in this run, replacing any prior report."""
    wb = Workbook()
    active_sheet = wb.active
    has_sheet = False
    if rows:
        headers = ["销售类型", "销售单号", "销售日期", "收款日期", "物料名称", "回收单号", "总金额", "税金", "总重量", "备注"]
        data = [[row.sale_type, row.sale_no, row.sale_date, row.payment_date, row.material,
                 row.recovery_no, row.total_amount, row.tax, row.total_weight, row.remark] for row in rows]
        write_sheet(active_sheet, "销售数据", headers, data, [12, 20, 14, 14, 22, 22, 14, 14, 14, 30])
        for row in active_sheet.iter_rows(min_row=2):
            row[4].alignment = row[5].alignment = row[9].alignment = Alignment(wrap_text=True, vertical="top")
        has_sheet = True
    for name, (headers, data) in sheets.items():
        ws = wb.create_sheet() if has_sheet else active_sheet
        write_sheet(ws, name, headers, data)
        has_sheet = True
    wb.save(path)


def collect_voucher_response_rows(page: Page, start: str, end: str, log, form_name: str,
                                  open_form) -> list[dict]:
    """Collect vc_bill_list responses and always detach the page listener."""
    response_rows: list[dict] = []
    response_count = 0
    total_pages = 0
    total_rows = 0
    page_size = 0
    capture_enabled = False

    def capture_response(response):
        nonlocal response_count, total_pages, total_rows, page_size
        if not capture_enabled:
            return
        if "erp.bfcgj.com" not in response.url:
            return
        try:
            payload = response.json()
        except Exception:
            return
        blocks = find_named_blocks(payload, "vc_bill_list")
        if blocks:
            for block in blocks:
                rows = list(block["rowsData"])
                response_rows.extend(rows)
                total_pages = max(total_pages, int(block.get("pages") or 0))
                total_rows = max(total_rows, int(block.get("totalRows") or 0))
                page_size = max(page_size, int(block.get("pageRows") or len(rows)))
            if not total_pages and total_rows and page_size:
                total_pages = (total_rows + page_size - 1) // page_size
            response_count += 1

    def wait_for_response(previous_count: int, timeout: float = 10, phase: str = "翻页") -> None:
        deadline = time.monotonic() + timeout
        while response_count <= previous_count and time.monotonic() < deadline:
            page.wait_for_timeout(200)
        if response_count <= previous_count:
            raise RuntimeError(f"{form_name}: {phase}未捕获 vc_bill_list 响应")

    def begin_filtered_capture():
        nonlocal capture_enabled, response_count, total_pages, total_rows, page_size
        response_rows.clear()
        response_count = 0
        total_pages = 0
        total_rows = 0
        page_size = 0
        capture_enabled = True

    page.on("response", capture_response)
    try:
        # 打开页面后先清空默认列表，再提交真实筛选请求。
        capture_enabled = True
        filter_table = open_form()
        page.wait_for_timeout(300)
        default_count = len(response_rows)
        if not apply_voucher_filters(page, filter_table, start, end, begin_filtered_capture):
            raise RuntimeError(f"{form_name}: 未识别到开始/结束日期或包含已入库筛选条件")
        log(f"{form_name}: 已提交筛选，日期 {start} 至 {end}，包含已入库")
        if default_count:
            log(f"{form_name}: 已忽略打开页面时的默认列表 {default_count} 条")
        wait_for_response(0, phase="列表查询")
        page_no = 1
        while True:
            log(f"{form_name}: 正在读取第 {page_no} 页")
            if total_pages and page_no >= total_pages:
                log(f"{form_name}: 接口共 {total_pages} 页")
                break
            # A stale ERP pager can remain visible after a form switch. If
            # the response did not provide page metadata, do not click it and
            # wait for a response that may belong to another view.
            if not total_pages:
                log(f"{form_name}: 接口未提供可靠分页信息，已停止在第 {page_no} 页")
                break
            next_button = page.locator('button[paging-btn="nextPage"]:visible').last
            if not next_button.count() or next_button.is_disabled() or "disabled" in (next_button.get_attribute("class") or ""):
                break
            before_page = response_count
            next_button.click()
            wait_for_response(before_page, phase="翻页")
            page_no += 1
        # The PDF pass runs after this listener is detached; preserve the
        # server-reported page count for that second pass.
        filtered_rows = list(response_rows)
        log(f"{form_name}: 日期筛选接口读取 {len(filtered_rows)} 条")
        setattr(page, "_facas_voucher_response_seen", bool(response_count))
        setattr(page, "_facas_voucher_total_pages", total_pages)
        return filtered_rows
    finally:
        remove_response_listener(page, capture_response)


def click_voucher_page_and_wait(page: Page, button, form_name: str, timeout: float = 10) -> None:
    """Click a voucher pager control and wait for its vc_bill_list response."""
    response_count = 0

    def capture_response(response):
        nonlocal response_count
        if "erp.bfcgj.com" not in response.url:
            return
        try:
            payload = response.json()
        except Exception:
            return
        if find_named_blocks(payload, "vc_bill_list"):
            response_count += 1

    page.on("response", capture_response)
    try:
        button.click(timeout=10000, force=True)
        deadline = time.monotonic() + timeout
        while response_count == 0 and time.monotonic() < deadline:
            page.wait_for_timeout(200)
        if response_count == 0:
            raise RuntimeError(f"{form_name} 翻页后未捕获 vc_bill_list 响应")
    finally:
        remove_response_listener(page, capture_response)


def scrape_default_form(page: Page, form_name: str, start: str, end: str, log,
                        pdf_dir: Path | None = None, generated_pdfs: set[Path] | None = None,
                        pdf_a5_landscape: bool = False) -> tuple[list[str], list[list[str]]]:
    """Read a voucher list using the columns rendered by the ERP itself."""
    # Voucher menus are nested under 财务管理 > 凭证处理 > 凭证探测器.
    category = "采购" if form_name == "采购单" else "车辆"
    menu_path = (
        ("财务管理", "凭证处理"),
        ("凭证处理", "凭证探测器"),
        ("凭证探测器", category),
        (category, form_name),
    )
    for parent_name, child_name in menu_path:
        expand_menu_section(page, parent_name, child_name)
    log(f"{form_name}: 正在打开凭证列表")
    def open_form():
        # The response listener must be installed before opening the form;
        # ERP commonly sends the initial vc_bill_list request immediately.
        click_visible_menu_item(page, form_name)
        page.wait_for_timeout(700)
        table = data_grid(page, ("票据号码", "单据日期"))
        table.wait_for(state="visible", timeout=15000)
        # 凭证探测器的日期条件默认隐藏，必须先打开过滤面板并确认。
        filter_button = page.get_by_role("button", name="过滤").last
        if not filter_button.count() or not filter_button.is_visible():
            filter_button = page.get_by_text("过滤", exact=True).last
        if not filter_button.count() or not filter_button.is_visible():
            raise RuntimeError(f"{form_name}: 未找到过滤按钮")
        filter_button.click(timeout=5000, force=True)
        page.wait_for_timeout(300)
        page.wait_for_timeout(300)
        filter_tables = page.locator("table.yc-view-free-table:visible")
        for index in range(filter_tables.count()):
            candidate = filter_tables.nth(index)
            if candidate.locator("input[filter-name='billDateFrom']:visible, input[filter-name='bill_datefrom']:visible").count():
                return candidate
        raise RuntimeError(f"{form_name}: 打开过滤面板后未找到日期筛选表格")

    response_rows = collect_voucher_response_rows(page, start, end, log, form_name, open_form)
    voucher_total_pages = getattr(page, "_facas_voucher_total_pages", 0)
    voucher_response_seen = getattr(page, "_facas_voucher_response_seen", False)
    log(f"{form_name}: 已获取 {len(response_rows)} 条接口记录")

    # Confirmed visible order, mapped to the vc_bill_list response fields.
    headers: list[str] = [
        "票据号码", "单据日期", "对方单位", "结算方式",
        "完成日期", "金额", "制单",
    ]
    rows_out: list[list[str]] = []
    if not response_rows and voucher_response_seen:
        log(f"{form_name}: 日期范围内未查询到凭证")
        return headers, []
    if not response_rows:
        raise RuntimeError("未捕获 vc_bill_list 响应，已停止导出以避免生成错误列数据")
    seen_response = set()
    for row in response_rows:
        bill_no = clean(str(row.get("c0", "")))
        if bill_no and bill_no in seen_response:
            continue
        if bill_no:
            seen_response.add(bill_no)
        values = [voucher_value(row, header) for header in headers]
        if any(values) and values not in rows_out:
            rows_out.append(values)
    if pdf_dir is not None:
        pending = {}
        for row in response_rows:
            bill_no = clean(str(row.get("c0", "")))
            if bill_no:
                pending.setdefault(bill_no, row)
        log(f"{form_name}: 开始批量保存 PDF，共 {len(pending)} 张")

        # List extraction ends on the final page. Return to the beginning and
        # process the visible bills page by page so earlier pages are not lost.
        previous_button = page.locator('button[paging-btn="prevPage"]:visible').last
        selected_page = "1"
        selected_items = page.locator("li.paging-item.paging-item-selected:visible")
        if selected_items.count():
            selected_page = clean(selected_items.last.inner_text())
        if selected_page != "1":
            for _ in range(200):
                if (not previous_button.count() or previous_button.is_disabled()
                        or "disabled" in (previous_button.get_attribute("class") or "")):
                    break
                click_voucher_page_and_wait(page, previous_button, form_name)

        while pending:
            grid = data_grid(page, ("票据号码", "单据日期"))
            for bill_no, row in list(pending.items()):
                log(f"{form_name}: 正在打开单据 {bill_no} 详情")
                # The API uses the internal number ($AF...), while the grid
                # normally renders only c1 (for example 00000080). Match
                # either representation instead of assuming both are shown.
                short_bill_no = clean(str(row.get("c1") or ""))
                bill_tokens = [token for token in (bill_no, short_bill_no) if token]
                matching_rows = []
                for row_index in range(grid.locator("tr").count()):
                    candidate_row = grid.locator("tr").nth(row_index)
                    try:
                        if not candidate_row.is_visible():
                            continue
                        row_text = clean(candidate_row.inner_text())
                        if any(token in row_text for token in bill_tokens):
                            matching_rows.append((candidate_row, row_text))
                    except PlaywrightError:
                        continue
                opened = False
                for candidate, row_text in matching_rows:
                    if candidate.is_visible():
                        try:
                            log(f"{form_name}: 已匹配票据行 {row_text[:120]}")
                            candidate.dblclick(timeout=5000)
                            opened = True
                            break
                        except PlaywrightError:
                            continue
                if not opened:
                    log(f"{form_name}: 当前页未找到单据 {bill_no}，保留待处理")
                    continue
                pending.pop(bill_no)
                log(f"{form_name}: 已打开单据 {bill_no} 详情")
                try:
                    page.wait_for_timeout(700)
                    bill_date = voucher_value(row, "单据日期") or start
                    pdf_path = pdf_dir / f"{safe_filename(form_name)}-{safe_filename(bill_date)}-{safe_filename(bill_no)}.pdf"
                    expected_print_values = [bill_no, short_bill_no, bill_date]
                    if save_printbill_pdf(
                            page, pdf_path, log, expected_print_values,
                            a5_landscape=pdf_a5_landscape):
                        if generated_pdfs is not None:
                            generated_pdfs.add(pdf_path.resolve())
                        log(f"已保存 PDF: {pdf_path.name}")
                    else:
                        log(f"未生成票据 PDF: {pdf_path.name}")
                except (PlaywrightError, PlaywrightTimeoutError) as exc:
                    log(f"{form_name} 单据 PDF 失败 ({bill_no}): {exc}")
                finally:
                    close = page.get_by_role("button", name="关闭").last
                    try:
                        if close.count() and close.is_visible():
                            close.click(timeout=3000, force=True, no_wait_after=True)
                            if not page.is_closed():
                                page.wait_for_timeout(300)
                    except PlaywrightError:
                        log(f"{form_name} 详情关闭失败: {bill_no}")
            if voucher_total_pages:
                selected_items = page.locator("li.paging-item.paging-item-selected:visible")
                selected_page = clean(selected_items.last.inner_text()) if selected_items.count() else ""
                if selected_page and selected_page.isdigit() and int(selected_page) >= voucher_total_pages:
                    break
            next_button = page.locator('button[paging-btn="nextPage"]:visible').last
            if not next_button.count() or next_button.is_disabled() or "disabled" in (next_button.get_attribute("class") or ""):
                break
            click_voucher_page_and_wait(page, next_button, form_name)
        for bill_no in pending:
            log(f"{form_name} 未找到可双击的票据行: {bill_no}")
    width = len(headers)
    rows_out = [row[:width] + [""] * max(0, width - len(row)) for row in rows_out]
    return headers, rows_out


def run(start: str, end: str, output: Path, selected: list[str], log,
        pdf_dir: Path | None = None, selected_forms: list[str] | None = None,
        save_excel: bool = True, pdf_a5_landscape: bool = False) -> tuple[int, int, bool]:
    if save_excel:
        output.parent.mkdir(parents=True, exist_ok=True)
    if pdf_dir is not None:
        pdf_dir.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        if not debug_port_available():
            start_debug_edge()
            raise LoginRequired("已启动专用 Edge。请在新窗口登录 ERP，登录完成后再次点击“开始提取”。")
        try:
            browser = p.chromium.connect_over_cdp(f"http://127.0.0.1:{DEBUG_PORT}")
        except PlaywrightError as exc:
            raise RuntimeError(f"无法连接专用 Edge: {exc}") from exc
        pages = [page for context in browser.contexts for page in context.pages]
        if not pages:
            raise RuntimeError("未找到已连接的 Edge 页面")
        page = live_erp_page(browser)
        rows: list[SaleRow] = []
        form_sheets: dict[str, tuple[list[str], list[list[str]]]] = {}
        generated_pdfs: set[Path] = set()
        for sale_type in selected:
            page = live_erp_page(browser)
            log(f"正在读取: {sale_type}")
            module_rows = scrape_module(
                page, sale_type, start, end, log, pdf_dir, generated_pdfs,
                pdf_a5_landscape)
            log(f"{sale_type}: 查询到 {len(module_rows)} 张销售单")
            rows.extend(module_rows)
        if not rows and not selected_forms:
            log("所选日期范围内未查询到可导出的销售单")
        if selected_forms:
            form_errors = []
            for form_name in selected_forms:
                page = live_erp_page(browser)
                log(f"正在读取: {form_name}")
                try:
                    form_sheets[form_name] = scrape_default_form(
                        page, form_name, start, end, log, pdf_dir, generated_pdfs,
                        pdf_a5_landscape)
                    log(f"{form_name}: 读取 {len(form_sheets[form_name][1])} 条")
                except Exception as exc:
                    form_errors.append(f"{form_name}: {exc}")
                    log(f"{form_name} 读取失败: {exc}")
            if not form_sheets and not rows:
                raise RuntimeError("凭证类型均未读取成功: " + "；".join(form_errors))
            for error in form_errors:
                log(f"跳过失败类型: {error}")
        if save_excel and (rows or form_sheets):
            log(f"正在生成 Excel: {output}")
            export_report(rows, form_sheets, output)
            log(f"Excel 已生成: {output}")
        excel_generated = bool(save_excel and (rows or form_sheets) and output.exists())
        form_row_count = sum(len(rows_data) for _, rows_data in form_sheets.values()) if selected_forms else 0
        total_count = len(rows) + form_row_count
        pdf_count = len(generated_pdfs)
        log(f"完成，共读取 {total_count} 条（销售 {len(rows)} 条，凭证 {form_row_count} 条）；实际生成 {pdf_count} 个 PDF")
        return total_count, pdf_count, excel_generated


class ExtractionWorker(QObject):
    """在 Qt 工作线程中执行 ERP 抓取，避免阻塞主界面。"""

    log_message = Signal(str)
    finished = Signal(int, int, bool)
    failed = Signal(str, str)

    def __init__(self, start: str, end: str, output: Path, selected: list[str],
                 pdf_dir: Path | None, selected_forms: list[str] | None,
                 save_excel: bool, pdf_a5_landscape: bool):
        super().__init__()
        self.start = start
        self.end = end
        self.output = output
        self.selected = selected
        self.pdf_dir = pdf_dir
        self.selected_forms = selected_forms
        self.save_excel = save_excel
        self.pdf_a5_landscape = pdf_a5_landscape

    def log(self, text: str) -> None:
        self.log_message.emit(str(text))

    @Slot()
    def execute(self) -> None:
        try:
            row_count, pdf_count, excel_generated = run(
                self.start,
                self.end,
                self.output,
                self.selected,
                self.log,
                self.pdf_dir,
                self.selected_forms,
                self.save_excel,
                self.pdf_a5_landscape,
            )
            outputs = [
                "日期范围内未查询到可导出数据"
                if row_count == 0 else f"共读取 {row_count} 条数据"
            ]
            outputs.append("已生成 1 个 Excel" if excel_generated else "未生成 Excel")
            outputs.append(
                f"已生成 {pdf_count} 个 PDF"
                if self.pdf_dir is not None else "未生成 PDF"
            )
            message = "；".join(outputs) + "。"
            self.log(f"完成: {message}")
            self.log(
                f"统计: 实际读取 {row_count} 条，Excel "
                f"{'1' if excel_generated else '0'} 个，PDF "
                f"{pdf_count if self.pdf_dir is not None else '0'} 个"
            )
            if excel_generated:
                self.log(f"Excel 文件: {self.output}")
            if self.pdf_dir is not None:
                self.log(f"PDF 目录: {self.pdf_dir}")
            self.finished.emit(row_count, pdf_count, excel_generated)
        except LoginRequired as exc:
            message = str(exc)
            self.log(message)
            self.failed.emit("login", message)
        except Exception as exc:
            error_text = str(exc)
            self.log(f"失败: {error_text}")
            self.log(f"错误类型: {type(exc).__name__}")
            self.log(traceback.format_exc().strip())
            self.failed.emit("error", error_text)


class ErpStateProbe(QObject):
    """在后台线程检测专用 Edge，避免端口探测阻塞主界面。"""

    state_ready = Signal(bool)

    def __init__(self):
        super().__init__()
        self._checking = False

    @Slot()
    def check(self) -> None:
        # 定时信号可能在上一次探测尚未结束时再次到达，避免并发探测。
        if self._checking:
            return
        self._checking = True
        try:
            current = debug_port_available()
        except OSError:
            current = False
        finally:
            self._checking = False
        self.state_ready.emit(current)


MODERN_STYLE = """
QMainWindow, QWidget#root {
    background: #edf4fc;
    color: #17324d;
}
QFrame#headerCard {
    background: #0a2948;
    border: 1px solid #164d78;
    border-radius: 18px;
}
QFrame#accentBar {
    background: #44c5f4;
    border-radius: 3px;
}
QLabel#eyebrow {
    color: #8fc7e8;
    font-size: 10px;
    font-weight: 700;
    letter-spacing: 1px;
}
QLabel#pageTitle {
    color: #ffffff;
    font-size: 22px;
    font-weight: 700;
}
QLabel#pageSubtitle {
    color: #b9d5ea;
    font-size: 11px;
}
QLabel#versionBadge {
    background: #123e65;
    color: #d9f4ff;
    border: 1px solid #2b719f;
    border-radius: 9px;
    padding: 5px 10px;
    font-size: 10px;
    font-weight: 700;
}
QLabel#statusChip {
    border-radius: 9px;
    padding: 5px 10px;
    font-size: 10px;
    font-weight: 700;
}
QLabel#statusChip[state="ready"] {
    background: #e3f2ff;
    color: #1261a0;
}
QLabel#statusChip[state="running"] {
    background: #e4f7ff;
    color: #087eaa;
}
QLabel#statusChip[state="success"] {
    background: #e1f7ef;
    color: #18794e;
}
QLabel#statusChip[state="error"] {
    background: #fde8ef;
    color: #b4235a;
}
QFrame#card {
    background: #ffffff;
    border: 1px solid #cfe0f2;
    border-radius: 14px;
}
QFrame#actionBar {
    background: #f7fbff;
    border: 1px solid #c5dced;
    border-radius: 14px;
}
QLabel#taskStatus {
    background: #e3f2ff;
    color: #1261a0;
    border: 1px solid #b7d8ef;
    border-radius: 9px;
    padding: 7px 12px;
    min-height: 18px;
    font-size: 11px;
    font-weight: 600;
}
QLabel#taskStatus[state="ready"] {
    background: #e3f2ff;
    color: #1261a0;
    border-color: #b7d8ef;
}
QLabel#taskStatus[state="running"] {
    background: #e4f7ff;
    color: #087eaa;
    border-color: #9ed8ec;
}
QLabel#taskStatus[state="success"] {
    background: #e1f7ef;
    color: #18794e;
    border-color: #a9dfc8;
}
QLabel#taskStatus[state="error"] {
    background: #fde8ef;
    color: #b4235a;
    border-color: #efb8ca;
}
QFrame#subCard {
    background: #f4f8fd;
    border: 1px solid #d8e6f5;
    border-radius: 10px;
}
QLabel#cardTitle {
    color: #17324d;
    font-size: 12px;
    font-weight: 700;
}
QLabel#cardHint {
    color: #6e8499;
    font-size: 10px;
}
QLabel#subTitle {
    color: #2e587b;
    font-size: 10px;
    font-weight: 700;
}
QLabel#fieldLabel {
    color: #54708a;
    font-size: 10px;
    font-weight: 600;
}
QLineEdit {
    background: #ffffff;
    border: 1px solid #bcd2e8;
    border-radius: 10px;
    padding: 7px 10px;
    min-height: 21px;
    selection-background-color: #b8e4fb;
}
QDateEdit {
    background: #ffffff;
    border: 1px solid #bcd2e8;
    border-radius: 10px;
    padding: 7px 6px 7px 11px;
    min-height: 21px;
    color: #17324d;
    font-size: 11px;
    selection-background-color: #b8e4fb;
}
QLineEdit:focus, QDateEdit:focus {
    border: 1px solid #2589c7;
}
QLineEdit:disabled, QDateEdit:disabled {
    background: #eaf1f8;
    color: #91a5b8;
}
QDateEdit QLineEdit {
    background: transparent;
    border: none;
    border-radius: 0;
    padding: 0 2px 0 0;
    min-height: 0;
    color: #17324d;
    font-size: 11px;
}
QDateEdit::drop-down {
    subcontrol-origin: padding;
    subcontrol-position: top right;
    width: 29px;
    border: none;
    margin: 1px;
    border-top-right-radius: 9px;
    border-bottom-right-radius: 9px;
    background: transparent;
}
QDateEdit::drop-down:hover {
    background: #e0f1fc;
}
QCheckBox {
    color: #294863;
    spacing: 8px;
    padding: 3px 0;
}
QCheckBox:checked {
    color: #1261a0;
    font-weight: 700;
}
QCheckBox::indicator {
    width: 17px;
    height: 17px;
    border: 1px solid #a9c4dc;
    border-radius: 5px;
    background: #ffffff;
}
QCheckBox::indicator:hover {
    border: 1px solid #2589c7;
}
QCheckBox::indicator:checked {
    background: #1676b8;
    border: 1px solid #1676b8;
}
QPushButton {
    border: 1px solid #bcd2e8;
    border-radius: 8px;
    background: #ffffff;
    color: #2e587b;
    padding: 7px 12px;
    font-weight: 600;
}
QPushButton:hover {
    background: #edf7ff;
    border-color: #75b8df;
}
QPushButton:pressed {
    background: #d9effc;
}
QPushButton:disabled {
    background: #eaf1f8;
    color: #91a5b8;
    border-color: #d9e5f0;
}
QPushButton#primaryButton {
    background: #1267a8;
    color: #ffffff;
    border: none;
    border-radius: 10px;
    padding: 10px 24px;
    font-size: 11px;
    font-weight: 700;
}
QPushButton#primaryButton:hover {
    background: #0e5590;
}
QPushButton#primaryButton:pressed {
    background: #0a416e;
}
QPushButton#primaryButton:disabled {
    background: #86b8d7;
    color: #eaf7ff;
}
QPushButton#secondaryButton {
    background: #e5f2fc;
    color: #1261a0;
    border: 1px solid #9fc8e5;
    border-radius: 10px;
    padding: 10px 18px;
    font-size: 11px;
    font-weight: 700;
}
QPushButton#secondaryButton:hover {
    background: #d7edfb;
    border-color: #62a9d2;
}
QPushButton#secondaryButton:pressed {
    background: #c6e5f7;
}
QPushButton#secondaryButton:disabled {
    background: #eaf1f8;
    color: #91a5b8;
    border-color: #d9e5f0;
}
QPushButton#quietButton {
    background: #edf5fc;
    border: 1px solid #d0e2f3;
    color: #3d6687;
    padding: 6px 10px;
    font-size: 10px;
}
QPlainTextEdit#logView {
    background: #0c2137;
    color: #c9eeff;
    border: 1px solid #17496d;
    border-radius: 10px;
    padding: 12px;
    selection-background-color: #15527b;
}
QScrollBar:vertical {
    background: #e6f0fa;
    width: 12px;
    margin: 0px;
    border: 1px solid #c7dceb;
    border-radius: 6px;
}
QScrollBar::sub-line:vertical,
QScrollBar::add-line:vertical {
    background: #dcebf7;
    height: 14px;
    subcontrol-origin: margin;
    border: 1px solid #c7dceb;
}
QScrollBar::sub-line:vertical {
    subcontrol-position: top;
    border-bottom: none;
    border-top-left-radius: 5px;
    border-top-right-radius: 5px;
}
QScrollBar::add-line:vertical {
    subcontrol-position: bottom;
    border-top: none;
    border-bottom-left-radius: 5px;
    border-bottom-right-radius: 5px;
}
QScrollBar::up-arrow:vertical,
QScrollBar::down-arrow:vertical {
    width: 8px;
    height: 8px;
    background: transparent;
}
QScrollBar::up-arrow:vertical {
    border-left: 4px solid transparent;
    border-right: 4px solid transparent;
    border-bottom: 5px solid #4f7fa2;
}
QScrollBar::down-arrow:vertical {
    border-left: 4px solid transparent;
    border-right: 4px solid transparent;
    border-top: 5px solid #4f7fa2;
}
QScrollBar::sub-line:vertical:hover,
QScrollBar::add-line:vertical:hover {
    background: #cce4f4;
}
QScrollBar::handle:vertical {
    background: #8cb9d8;
    min-height: 30px;
    margin: 1px;
    border-radius: 4px;
}
QScrollBar::handle:vertical:hover {
    background: #5d9bc4;
}
QProgressBar {
    border: none;
    background: #dcebf8;
    border-radius: 3px;
    min-height: 6px;
    max-height: 6px;
}
QProgressBar::chunk {
    background: #27a8df;
    border-radius: 3px;
}
"""


DATE_CALENDAR_STYLE = """
QCalendarWidget {
    background: #ffffff;
    color: #173b5f;
    border: 1px solid #9fc4e2;
    border-radius: 12px;
}
QCalendarWidget QWidget#qt_calendar_navigationbar {
    background: #eaf4ff;
    border-bottom: 1px solid #c9dff2;
    border-top-left-radius: 11px;
    border-top-right-radius: 11px;
    padding: 5px;
}
QCalendarWidget QToolButton {
    color: #174d7d;
    background: transparent;
    border: none;
    border-radius: 6px;
    padding: 5px 8px;
    font-weight: 700;
}
QCalendarWidget QToolButton:hover {
    background: #dbeeff;
    color: #0b63ae;
}
QCalendarWidget QToolButton:pressed {
    background: #c7e5fb;
}
QCalendarWidget QSpinBox {
    color: #173b5f;
    background: #ffffff;
    border: 1px solid #b8d4eb;
    border-radius: 5px;
    padding: 2px 5px;
}
QCalendarWidget QSpinBox::up-button,
QCalendarWidget QSpinBox::down-button {
    width: 14px;
    border: none;
    background: transparent;
}
QCalendarWidget QAbstractItemView {
    background: #ffffff;
    alternate-background-color: #eef6fd;
    color: #244967;
    selection-background-color: #2184c1;
    selection-color: #ffffff;
    outline: none;
    border: none;
    padding: 6px;
}
QCalendarWidget QAbstractItemView::item:selected {
    background: #2184c1;
    color: #ffffff;
    border-radius: 8px;
}
QCalendarWidget QAbstractItemView::item:hover {
    background: #d9effc;
    color: #0d4d7d;
}
QCalendarWidget QMenu {
    background: #ffffff;
    color: #173b5f;
    border: 1px solid #b8d4eb;
}
"""


class App(QMainWindow):
    """PySide6 主界面，负责交互、配置和后台任务调度。"""

    erp_probe_requested = Signal()

    def __init__(self):
        super().__init__()
        self.setObjectName("mainWindow")
        self.setWindowTitle(f"报废汽车财务数据自动化处理 {APP_VERSION}")
        # 固定尺寸，避免复选项、输出路径和日志区在缩放时重新换行。
        self.setFixedSize(1250, 770)
        self.setStyleSheet(MODERN_STYLE)

        self.app_dir = application_dir()
        output_root = self.app_dir / "output"
        self.default_output = output_root / "Excel" / "销售数据.xlsx"
        self.default_pdf_dir = output_root / "PDF"
        self.settings_path = self.app_dir / "facas-settings.json"
        self.log_dir = self.app_dir / "log"
        self.running = False
        self.bank_task_running = False
        self.erp_open = False
        self.erp_launching = False
        self.bank_window: BankEntryWindow | None = None
        self.thread: QThread | None = None
        self.worker: ExtractionWorker | None = None

        self._build_ui()
        self.load_settings()
        self._update_output_controls()
        self._update_erp_controls()

        # 端口探测放到独立线程，避免 ERP 未启动时连接超时卡住 Qt 主线程。
        self.erp_probe_thread = QThread(self)
        self.erp_probe_worker = ErpStateProbe()
        self.erp_probe_worker.moveToThread(self.erp_probe_thread)
        self.erp_probe_requested.connect(self.erp_probe_worker.check, Qt.ConnectionType.QueuedConnection)
        self.erp_probe_worker.state_ready.connect(self._apply_erp_state, Qt.ConnectionType.QueuedConnection)
        self.erp_probe_thread.finished.connect(self.erp_probe_worker.deleteLater)
        self.erp_probe_thread.start()

        self._refresh_erp_state()
        self.erp_state_timer = QTimer(self)
        self.erp_state_timer.setInterval(1000)
        self.erp_state_timer.timeout.connect(self._refresh_erp_state)
        self.erp_state_timer.start()
        self._center_window()
        self.log(f"程序已启动，日志文件: {self.log_dir / f'{date.today().isoformat()}.log'}")

    def _card(self, title: str, hint: str = "") -> tuple[QFrame, QVBoxLayout]:
        card = QFrame()
        card.setObjectName("card")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(16, 14, 16, 16)
        layout.setSpacing(10)
        title_row = QHBoxLayout()
        title_label = QLabel(title)
        title_label.setObjectName("cardTitle")
        title_row.addWidget(title_label)
        title_row.addStretch(1)
        if hint:
            hint_label = QLabel(hint)
            hint_label.setObjectName("cardHint")
            title_row.addWidget(hint_label)
        layout.addLayout(title_row)
        return card, layout

    @staticmethod
    def _field_label(text: str) -> QLabel:
        label = QLabel(text)
        label.setObjectName("fieldLabel")
        return label

    @staticmethod
    def _configure_date_edit(control: QDateEdit) -> None:
        """统一日期输入框字体、尺寸、间距与蓝色主题。"""
        control.setCalendarPopup(True)
        control.setDisplayFormat("yyyy-MM-dd")
        control.setFixedWidth(184)
        control.setMinimumHeight(40)
        control.setFrame(True)
        # 部分 Windows 电脑的系统字体或缩放比例不同，显式同步编辑框字体并留足显示宽度。
        control.setFont(QApplication.font())
        editor = control.lineEdit()
        editor.setFont(QApplication.font())
        editor.setAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        calendar = control.calendarWidget()
        calendar.setObjectName("dateCalendar")
        calendar.setStyleSheet(DATE_CALENDAR_STYLE)
        calendar.setGridVisible(False)
        calendar.setVerticalHeaderFormat(QCalendarWidget.VerticalHeaderFormat.NoVerticalHeader)
        calendar.setFirstDayOfWeek(Qt.DayOfWeek.Monday)
        calendar.setMinimumSize(318, 248)
        calendar.setFont(QFont("Microsoft YaHei UI", 10))

    def _build_ui(self) -> None:
        root = QWidget()
        root.setObjectName("root")
        self.setCentralWidget(root)
        outer = QVBoxLayout(root)
        outer.setContentsMargins(26, 24, 26, 22)
        outer.setSpacing(16)

        header = QFrame()
        header.setObjectName("headerCard")
        header_layout = QHBoxLayout(header)
        header_layout.setContentsMargins(20, 18, 20, 18)
        header_layout.setSpacing(14)
        accent = QFrame()
        accent.setObjectName("accentBar")
        accent.setFixedWidth(5)
        header_layout.addWidget(accent)
        title_box = QVBoxLayout()
        title_box.setSpacing(4)
        eyebrow = QLabel("FACAS  /  INTERNAL AUTOMATION")
        eyebrow.setObjectName("eyebrow")
        title_box.addWidget(eyebrow)
        title = QLabel("财务数据自动化处理")
        title.setObjectName("pageTitle")
        title_box.addWidget(title)
        subtitle = QLabel("销售与凭证数据采集  ·  Excel / PDF 输出")
        subtitle.setObjectName("pageSubtitle")
        title_box.addWidget(subtitle)
        header_layout.addLayout(title_box)
        header_layout.addStretch(1)
        meta = QVBoxLayout()
        meta.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        version = QLabel(APP_VERSION)
        version.setObjectName("versionBadge")
        version.setAlignment(Qt.AlignmentFlag.AlignCenter)
        meta.addWidget(version, alignment=Qt.AlignmentFlag.AlignRight)
        self.status_chip = QLabel("就绪")
        self.status_chip.setObjectName("statusChip")
        self.status_chip.setProperty("state", "ready")
        self.status_chip.setAlignment(Qt.AlignmentFlag.AlignCenter)
        meta.addWidget(self.status_chip, alignment=Qt.AlignmentFlag.AlignRight)
        header_layout.addLayout(meta)
        outer.addWidget(header)

        content = QHBoxLayout()
        content.setSpacing(16)
        left = QWidget()
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(0, 0, 0, 0)
        left_layout.setSpacing(12)
        # 左侧使用固定宽度，保证复选项和输出路径两列在不同窗口尺寸下不跳动。
        left.setFixedWidth(560)
        left.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Expanding)

        query_card, query_layout = self._card("查询范围", "销售和凭证均按此日期执行")
        date_row = QHBoxLayout()
        date_row.setSpacing(8)
        date_row.addWidget(self._field_label("开始日期"))
        self.start_date = DateEdit(QDate.currentDate())
        self._configure_date_edit(self.start_date)
        date_row.addWidget(self.start_date)
        date_row.addSpacing(12)
        date_row.addWidget(self._field_label("结束日期"))
        self.end_date = DateEdit(QDate.currentDate())
        self._configure_date_edit(self.end_date)
        date_row.addWidget(self.end_date)
        date_row.addStretch(1)
        query_layout.addLayout(date_row)
        left_layout.addWidget(query_card)

        scope_card, scope_layout = self._card("数据范围", "可同时选择销售分类和凭证类型")
        scope_actions = QHBoxLayout()
        scope_actions.addStretch(1)
        scope_actions.setSpacing(10)
        self.select_all_button = QPushButton("全选")
        self.select_all_button.setObjectName("quietButton")
        self.select_all_button.clicked.connect(lambda: self.set_categories(True))
        self.clear_all_button = QPushButton("清空")
        self.clear_all_button.setObjectName("quietButton")
        self.clear_all_button.clicked.connect(lambda: self.set_categories(False))
        scope_actions.addWidget(self.select_all_button)
        scope_actions.addWidget(self.clear_all_button)
        scope_layout.addLayout(scope_actions)
        groups = QGridLayout()
        groups.setContentsMargins(0, 2, 0, 0)
        groups.setHorizontalSpacing(12)
        groups.setVerticalSpacing(10)
        self.category_checks: dict[str, QCheckBox] = {}
        self.form_checks: dict[str, QCheckBox] = {}
        sales_group = self._selection_group("销售管理", MODULES, self.category_checks, True)
        voucher_group = self._selection_group("凭证探测器", FORM_MODULES, self.form_checks, False)
        groups.addWidget(sales_group, 0, 0)
        groups.addWidget(voucher_group, 0, 1)
        groups.setColumnMinimumWidth(0, 248)
        groups.setColumnMinimumWidth(1, 248)
        groups.setColumnStretch(0, 1)
        groups.setColumnStretch(1, 1)
        scope_layout.addLayout(groups)
        left_layout.addWidget(scope_card)

        output_card, output_layout = self._card("输出设置", "文件只会在对应开关开启时生成")
        output_layout.setSpacing(8)
        excel_row = QHBoxLayout()
        excel_row.setSpacing(8)
        self.save_excel = QCheckBox("生成 Excel")
        self.save_excel.setChecked(True)
        self.save_excel.setFixedWidth(92)
        self.save_excel.setMinimumHeight(38)
        self.save_excel.toggled.connect(self._update_output_controls)
        excel_row.addWidget(self.save_excel)
        self.output_edit = QLineEdit(str(self.default_output))
        self.output_edit.setPlaceholderText("选择 Excel 输出文件")
        self.output_edit.setMinimumHeight(38)
        excel_row.addWidget(self.output_edit, 1)
        self.output_browse = QPushButton("选择文件")
        self.output_browse.setFixedWidth(82)
        self.output_browse.setMinimumHeight(38)
        self.output_browse.clicked.connect(self.choose_output)
        excel_row.addWidget(self.output_browse)
        output_layout.addLayout(excel_row)

        pdf_row = QHBoxLayout()
        pdf_row.setSpacing(8)
        self.save_pdfs = QCheckBox("生成 PDF")
        self.save_pdfs.setChecked(True)
        self.save_pdfs.setFixedWidth(92)
        self.save_pdfs.setMinimumHeight(38)
        self.save_pdfs.toggled.connect(self._update_output_controls)
        pdf_row.addWidget(self.save_pdfs)
        self.pdf_dir_edit = QLineEdit(str(self.default_pdf_dir))
        self.pdf_dir_edit.setPlaceholderText("选择 PDF 输出目录")
        self.pdf_dir_edit.setMinimumHeight(38)
        pdf_row.addWidget(self.pdf_dir_edit, 1)
        self.pdf_browse = QPushButton("选择目录")
        self.pdf_browse.setFixedWidth(82)
        self.pdf_browse.setMinimumHeight(38)
        self.pdf_browse.clicked.connect(self.choose_pdf_dir)
        pdf_row.addWidget(self.pdf_browse)
        output_layout.addLayout(pdf_row)

        pdf_options = QHBoxLayout()
        pdf_options.setContentsMargins(100, 0, 0, 0)
        self.pdf_a5_landscape = QCheckBox("A5 横版")
        self.pdf_a5_landscape.setChecked(False)
        pdf_options.addWidget(self.pdf_a5_landscape)
        pdf_options.addWidget(QLabel("未勾选时使用 A4 竖版"))
        pdf_options.addStretch(1)
        output_layout.addLayout(pdf_options)
        left_layout.addWidget(output_card)
        left_layout.addStretch(1)

        log_card, log_layout = self._card("运行日志", "按日期保存在 log 文件夹")
        log_toolbar = QHBoxLayout()
        log_toolbar.addStretch(1)
        self.clear_log_button = QPushButton("清空显示")
        self.clear_log_button.setObjectName("quietButton")
        self.clear_log_button.clicked.connect(self.clear_log_view)
        log_toolbar.addWidget(self.clear_log_button)
        log_layout.addLayout(log_toolbar)
        self.log_view = QPlainTextEdit()
        self.log_view.setObjectName("logView")
        self.log_view.setReadOnly(True)
        self.log_view.setVerticalScrollBar(LogScrollBar(Qt.Orientation.Vertical, self.log_view))
        self.log_view.setLineWrapMode(QPlainTextEdit.LineWrapMode.WidgetWidth)
        self.log_view.setMaximumBlockCount(2500)
        self.log_view.setFont(QFont("Cascadia Mono", 9))
        self.log_view.setMinimumHeight(150)
        log_layout.addWidget(self.log_view, 1)
        # 日志区使用左侧布局之外的全部剩余宽度，并与左侧内容区等高。
        log_card.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        content.addWidget(left, 0)
        content.addWidget(log_card, 1)
        outer.addLayout(content, 1)

        action_bar = QFrame()
        action_bar.setObjectName("actionBar")
        action_layout = QHBoxLayout(action_bar)
        action_layout.setContentsMargins(16, 10, 16, 10)
        self.task_status = QLabel("请先登录 ERP，再选择日期和数据范围开始提取")
        self.task_status.setObjectName("taskStatus")
        self.task_status.setProperty("state", "ready")
        self.task_status.setMinimumWidth(360)
        self.task_status.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        action_layout.addWidget(self.task_status)
        action_layout.addStretch(1)
        self.progress = QProgressBar()
        self.progress.setRange(0, 0)
        self.progress.setFixedWidth(150)
        self.progress.setVisible(False)
        action_layout.addWidget(self.progress)
        self.login_button = QPushButton("登录 ERP")
        self.login_button.setObjectName("secondaryButton")
        self.login_button.setMinimumWidth(120)
        self.login_button.clicked.connect(self.login_erp)
        action_layout.addWidget(self.login_button)
        self.bank_button = QPushButton("银行流水录入")
        self.bank_button.setObjectName("secondaryButton")
        self.bank_button.setMinimumWidth(126)
        self.bank_button.clicked.connect(self.open_bank_entry)
        action_layout.addWidget(self.bank_button)
        self.start_button = QPushButton("开始提取")
        self.start_button.setObjectName("primaryButton")
        self.start_button.setMinimumWidth(140)
        self.start_button.clicked.connect(self.start_run)
        action_layout.addWidget(self.start_button)
        outer.addWidget(action_bar)

        self.task_controls = [
            self.start_date,
            self.end_date,
            *self.category_checks.values(),
            *self.form_checks.values(),
            self.save_excel,
            self.output_edit,
            self.output_browse,
            self.save_pdfs,
            self.pdf_dir_edit,
            self.pdf_browse,
            self.pdf_a5_landscape,
            self.select_all_button,
            self.clear_all_button,
            self.login_button,
            self.bank_button,
            self.start_button,
        ]

    @staticmethod
    def _selection_group(title: str, values: dict[str, str], storage: dict[str, QCheckBox], default: bool) -> QFrame:
        group = QFrame()
        group.setObjectName("subCard")
        group.setMinimumWidth(248)
        group.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        layout = QVBoxLayout(group)
        layout.setContentsMargins(14, 10, 14, 10)
        layout.setSpacing(6)
        label = QLabel(title)
        label.setObjectName("subTitle")
        label.setMinimumHeight(22)
        layout.addWidget(label)
        for name in values:
            check = QCheckBox(name)
            check.setChecked(default)
            check.setMinimumWidth(132)
            check.setMinimumHeight(28)
            check.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
            check.setToolTip(name)
            storage[name] = check
            layout.addWidget(check)
        layout.addStretch(1)
        return group

    def _center_window(self) -> None:
        screen = self.screen()
        if screen is None:
            return
        available = screen.availableGeometry()
        self.move(available.center() - self.rect().center())

    def _update_output_controls(self) -> None:
        excel_enabled = self.save_excel.isChecked()
        pdf_enabled = self.save_pdfs.isChecked()
        busy = self.running or self.bank_task_running
        self.output_edit.setEnabled(excel_enabled and not busy)
        self.output_browse.setEnabled(excel_enabled and not busy)
        self.pdf_dir_edit.setEnabled(pdf_enabled and not busy)
        self.pdf_browse.setEnabled(pdf_enabled and not busy)
        self.pdf_a5_landscape.setEnabled(pdf_enabled and not busy)

    def open_bank_entry(self) -> None:
        """打开银行流水处理和 ERP 凭证录入界面。"""
        if self.bank_window is None:
            self.bank_window = BankEntryWindow(self)
            self.bank_window.bank_task_started.connect(self._bank_task_started)
            self.bank_window.bank_task_finished.connect(self._bank_task_finished)
        self.bank_window.set_external_busy(self.running or self.bank_task_running)
        self.bank_window.show()
        self.bank_window.raise_()
        self.bank_window.activateWindow()

    @Slot()
    def _bank_task_started(self) -> None:
        self.bank_task_running = True
        self._set_controls_enabled(False)
        self._set_status("运行中", "running", "银行流水任务正在执行…")
        self.log("银行流水任务已启动")

    @Slot(str)
    def _bank_task_finished(self, outcome: str) -> None:
        self.bank_task_running = False
        self._set_controls_enabled(True)
        if outcome == "success":
            self._set_status("已完成", "success", "银行流水任务已完成")
        elif outcome == "partial":
            self._set_status("部分失败", "error", "部分凭证未录入，请查看银行窗口日志")
        else:
            self._set_status("执行失败", "error", "银行流水任务失败，请查看银行窗口日志")
        self.log(f"银行流水任务结果：{outcome}")
        self._refresh_erp_state()

    def _update_erp_controls(self) -> None:
        """根据专用 Edge 状态统一锁定登录和提取按钮。"""
        if not hasattr(self, "login_button"):
            return
        if self.running or self.bank_task_running:
            self.login_button.setEnabled(False)
            self.start_button.setEnabled(False)
            self.bank_button.setEnabled(False)
            return
        self.bank_button.setEnabled(True)
        self.login_button.setEnabled(not self.erp_open and not self.erp_launching)
        self.start_button.setEnabled(self.erp_open)

    def _refresh_erp_state(self) -> None:
        """请求后台检测专用 Edge，不在界面线程中直接连接端口。"""
        if hasattr(self, "erp_probe_worker"):
            self.erp_probe_requested.emit()

    def _apply_erp_state(self, current: bool) -> None:
        """接收后台探测结果并更新界面控件状态。"""
        if current != self.erp_open:
            self.erp_open = current
            if current:
                self.erp_launching = False
                if not self.running and not self.bank_task_running:
                    self._set_status("ERP 已连接", "ready", "可以开始提取")
            elif not self.running and not self.bank_task_running:
                self.erp_launching = False
                self._set_status("待登录", "ready", "请先点击“登录 ERP”")
        self._update_erp_controls()

    def login_erp(self) -> None:
        """启动专用 Edge，让用户先完成 ERP 登录。"""
        if self.running or self.bank_task_running or self.erp_launching:
            return
        if self.erp_open:
            self.erp_open = True
            self._update_erp_controls()
            self.log("专用 Edge 已在运行，请完成 ERP 登录")
            self._set_status("ERP 已连接", "ready", "可以开始提取")
            return
        try:
            self.erp_launching = True
            self._update_erp_controls()
            start_debug_edge()
            self.log("已启动专用 Edge，请在新窗口完成 ERP 登录")
            self._set_status("待登录", "running", "请在专用 Edge 中完成 ERP 登录")
            QMessageBox.information(
                self,
                "登录 ERP",
                "专用 Edge 已启动，请在新窗口完成 ERP 登录。登录完成后点击“开始提取”。",
            )
        except Exception as exc:
            self.erp_launching = False
            self.log(f"启动专用 Edge 失败: {exc}")
            self._set_status("启动失败", "error", "请检查 Microsoft Edge 是否已安装")
            self._update_erp_controls()
            QMessageBox.critical(self, "无法启动 Edge", str(exc))

    def _set_controls_enabled(self, enabled: bool) -> None:
        for control in self.task_controls:
            control.setEnabled(enabled)
        if self.bank_window is not None:
            self.bank_window.set_external_busy(not enabled or self.bank_task_running)
        self._update_output_controls()
        self._update_erp_controls()

    def set_categories(self, enabled: bool) -> None:
        for check in [*self.category_checks.values(), *self.form_checks.values()]:
            check.setChecked(enabled)

    def choose_output(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self,
            "选择 Excel 输出文件",
            self.output_edit.text(),
            "Excel 文件 (*.xlsx)",
        )
        if path:
            if not path.lower().endswith(".xlsx"):
                path += ".xlsx"
            self.output_edit.setText(path)

    def choose_pdf_dir(self) -> None:
        path = QFileDialog.getExistingDirectory(self, "选择 PDF 输出目录", self.pdf_dir_edit.text())
        if path:
            self.pdf_dir_edit.setText(path)

    def load_settings(self) -> None:
        try:
            settings = json.loads(self.settings_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return
        for key, control in (("start", self.start_date), ("end", self.end_date)):
            value = settings.get(key)
            parsed = QDate.fromString(str(value), "yyyy-MM-dd") if value else QDate()
            if parsed.isValid():
                control.setDate(parsed)
        old_default_output = self.app_dir / "销售数据.xlsx"
        old_default_pdf_dir = self.app_dir / "销售单据PDF"
        stored_output = settings.get("output")
        if stored_output:
            stored_path = Path(stored_output)
            output = self.default_output if stored_path == old_default_output else stored_path
            self.output_edit.setText(str(output))
        stored_pdf_dir = settings.get("pdf_dir")
        if stored_pdf_dir:
            stored_path = Path(stored_pdf_dir)
            pdf_dir = self.default_pdf_dir if stored_path == old_default_pdf_dir else stored_path
            self.pdf_dir_edit.setText(str(pdf_dir))
        self.save_pdfs.setChecked(bool(settings.get("save_pdfs", True)))
        self.pdf_a5_landscape.setChecked(bool(settings.get("pdf_a5_landscape", False)))
        self.save_excel.setChecked(bool(settings.get("save_excel", True)))
        for name, check in self.category_checks.items():
            if name in settings.get("categories", {}):
                check.setChecked(bool(settings["categories"][name]))
        for name, check in self.form_checks.items():
            if name in settings.get("forms", {}):
                check.setChecked(bool(settings["forms"][name]))

    def save_settings(self) -> None:
        settings = {
            "start": self.start_date.date().toString("yyyy-MM-dd"),
            "end": self.end_date.date().toString("yyyy-MM-dd"),
            "output": self.output_edit.text(),
            "pdf_dir": self.pdf_dir_edit.text(),
            "save_pdfs": self.save_pdfs.isChecked(),
            "pdf_a5_landscape": self.pdf_a5_landscape.isChecked(),
            "save_excel": self.save_excel.isChecked(),
            "categories": {name: check.isChecked() for name, check in self.category_checks.items()},
            "forms": {name: check.isChecked() for name, check in self.form_checks.items()},
        }
        try:
            self.settings_path.write_text(
                json.dumps(settings, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        except OSError:
            pass

    def log(self, text: str) -> None:
        timestamped = f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {text}"
        try:
            self.log_dir.mkdir(parents=True, exist_ok=True)
            run_log_path = self.log_dir / f"{date.today().isoformat()}.log"
            with run_log_path.open("a", encoding="utf-8") as handle:
                handle.write(timestamped + "\n")
        except OSError:
            pass
        if hasattr(self, "log_view"):
            self.log_view.appendPlainText(timestamped)
            scrollbar = self.log_view.verticalScrollBar()
            scrollbar.setValue(scrollbar.maximum())

    def clear_log_view(self) -> None:
        self.log_view.clear()
        self.log_view.appendPlainText("界面日志已清空，完整记录仍保存在 log 文件夹。")

    def _set_status(self, text: str, state: str, detail: str | None = None) -> None:
        self.status_chip.setText(text)
        self.status_chip.setProperty("state", state)
        self.status_chip.style().unpolish(self.status_chip)
        self.status_chip.style().polish(self.status_chip)
        self.task_status.setProperty("state", state)
        self.task_status.style().unpolish(self.task_status)
        self.task_status.style().polish(self.task_status)
        if detail:
            self.task_status.setText(detail)

    def start_run(self) -> None:
        if self.running or self.bank_task_running:
            return
        selected = [name for name, check in self.category_checks.items() if check.isChecked()]
        selected_forms = [name for name, check in self.form_checks.items() if check.isChecked()]
        if not selected and not selected_forms:
            QMessageBox.warning(self, "提示", "至少选择一个销售分类或凭证类型")
            return
        start = self.start_date.date().toString("yyyy-MM-dd")
        end = self.end_date.date().toString("yyyy-MM-dd")
        if start > end:
            QMessageBox.warning(self, "日期范围错误", "开始日期不能晚于结束日期")
            return
        save_excel = self.save_excel.isChecked()
        save_pdfs = self.save_pdfs.isChecked()
        output_text = self.output_edit.text().strip()
        pdf_dir_text = self.pdf_dir_edit.text().strip()
        if save_excel and not output_text:
            QMessageBox.warning(self, "输出设置不完整", "请先选择 Excel 输出文件")
            return
        if save_pdfs and not pdf_dir_text:
            QMessageBox.warning(self, "输出设置不完整", "请先选择 PDF 输出目录")
            return
        if not self.erp_open:
            self.erp_open = False
            self.erp_launching = False
            self._update_erp_controls()
            message = "请先点击“登录 ERP”，在专用 Edge 中完成登录后再开始提取。"
            self.log(message)
            self._set_status("未登录", "error", message)
            QMessageBox.information(self, "请先登录 ERP", message)
            return

        self.save_settings()
        output = Path(output_text or self.default_output)
        pdf_dir = Path(pdf_dir_text) if save_pdfs else None
        pdf_a5_landscape = self.pdf_a5_landscape.isChecked()
        self.running = True
        self._set_controls_enabled(False)
        self.start_button.setText("正在提取…")
        self.progress.setVisible(True)
        self._set_status("运行中", "running", "正在连接 ERP 并执行任务…")
        self.log(f"任务已启动，日期范围: {start} 至 {end}")
        self.log(f"销售分类: {', '.join(selected) if selected else '未选择'}")
        self.log(f"凭证类型: {', '.join(selected_forms) if selected_forms else '未选择'}")
        pdf_status = "关闭"
        if pdf_dir is not None:
            pdf_status = f"开启（{'A5横版' if pdf_a5_landscape else 'A4竖版'}）"
        self.log(f"输出设置: Excel={'开启' if save_excel else '关闭'}；PDF={pdf_status}")
        if save_excel:
            self.log(f"Excel 路径: {output}")
        if pdf_dir is not None:
            self.log(f"PDF 目录: {pdf_dir}")

        self.thread = QThread(self)
        self.worker = ExtractionWorker(
            start,
            end,
            output,
            selected,
            pdf_dir,
            selected_forms,
            save_excel,
            pdf_a5_landscape,
        )
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.execute)
        self.worker.log_message.connect(self.log)
        self.worker.finished.connect(self._task_finished)
        self.worker.failed.connect(self._task_failed)
        self.worker.finished.connect(self.thread.quit)
        self.worker.failed.connect(self.thread.quit)
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.failed.connect(self.worker.deleteLater)
        self.thread.finished.connect(self._thread_finished)
        self.thread.finished.connect(self.thread.deleteLater)
        self.thread.start()

    @Slot(int, int, bool)
    def _task_finished(self, row_count: int, pdf_count: int, excel_generated: bool) -> None:
        outputs = [
            "日期范围内未查询到可导出数据"
            if row_count == 0 else f"共读取 {row_count} 条数据"
        ]
        outputs.append("已生成 1 个 Excel" if excel_generated else "未生成 Excel")
        outputs.append(
            f"已生成 {pdf_count} 个 PDF"
            if self.save_pdfs.isChecked() else "未生成 PDF"
        )
        message = "；".join(outputs) + "。"
        self._set_status("已完成", "success", message)
        if excel_generated:
            open_output_directory(Path(self.output_edit.text()))
        elif self.save_pdfs.isChecked() and pdf_count:
            open_output_directory(Path(self.pdf_dir_edit.text()))
        QMessageBox.information(self, "任务完成", message)

    @Slot(str, str)
    def _task_failed(self, kind: str, message: str) -> None:
        if kind == "login":
            self._set_status("需要登录", "running", "请在专用 Edge 中完成 ERP 登录")
            QMessageBox.information(self, "请先登录 ERP", message)
        else:
            self._set_status("执行失败", "error", message)
            QMessageBox.critical(self, "执行失败", message)

    @Slot()
    def _thread_finished(self) -> None:
        self.running = False
        self.erp_launching = False
        self.progress.setVisible(False)
        self.start_button.setText("开始提取")
        self._set_controls_enabled(True)
        self._refresh_erp_state()
        self.worker = None
        self.thread = None

    def closeEvent(self, event) -> None:
        if self.running or self.bank_task_running:
            QMessageBox.warning(self, "任务进行中", "当前任务仍在运行，请等待任务完成后再关闭窗口")
            event.ignore()
            return
        if hasattr(self, "erp_state_timer"):
            self.erp_state_timer.stop()
        if hasattr(self, "erp_probe_thread") and self.erp_probe_thread.isRunning():
            self.erp_probe_thread.quit()
            self.erp_probe_thread.wait(1000)
        self.save_settings()
        event.accept()


def qt_main() -> int:
    """启动 PySide6 主界面。"""
    qt_app = QApplication(sys.argv)
    qt_app.setApplicationName("FACAS")
    qt_app.setApplicationVersion(APP_VERSION)
    qt_app.setStyle("Fusion")
    qt_app.setFont(QFont("Microsoft YaHei UI", 10))
    window = App()
    window.show()
    return qt_app.exec()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--start"); parser.add_argument("--end"); parser.add_argument("--output", type=Path); parser.add_argument("--pdf-dir", type=Path)
    parser.add_argument("--a5-landscape", action="store_true", help="使用 A5 横版生成 PDF")
    args = parser.parse_args()
    if args.start and args.end:
        cli_output = args.output or (Path.cwd() / "output" / "Excel" / "销售数据.xlsx")
        run(
            args.start, args.end, cli_output, list(MODULES), print,
            pdf_dir=args.pdf_dir, pdf_a5_landscape=args.a5_landscape)
    else:
        raise SystemExit(qt_main())
